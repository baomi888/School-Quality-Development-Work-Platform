# -*- coding: utf-8 -*-
"""生成演示数据（录屏 / 自检用）——可随时重建，不含任何真实姓名。

用法：python make_demo_data.py [--out 目录]

产出一份干净的、能讲完整故事的演示素材：
  demo/收件箱/          5 份单表（不同学院各交一份，覆盖 5 种典型情况）
  demo/申报书_待预审.docx  一份故意写错的申报书，用于演示预审

设计要点：
  · 文件名为中性学院名，画面上不会暴露真实身份
  · 内容为虚构活动与虚构联系人（「示例·甲」这类占位名）
  · 每个文件对应一个要演示的判定，录屏时能一句话说清「这份是想让你看什么」
"""
import io
import shutil
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

HERE = Path(__file__).resolve().parent
# 默认输出到工程根下一级 _demo；可用 --out 指定
DEMO = Path(sys.argv[sys.argv.index('--out') + 1]).resolve() \
    if '--out' in sys.argv else HERE.parent.parent / '_demo'
INBOX = DEMO / '收件箱'
APPDOC = DEMO / '申报书_待预审.docx'

# 校区不硬编码：从工程配置里读，保证演示数据与当前部署的配置一致
sys.path.insert(0, str(HERE))
from core import zjb  # noqa: E402

_CFG = zjb.load_config(HERE / 'core' / 'config.json')
CAMPUSES = list(_CFG.get('campuses') or ['西城', '官渡'])
C1 = CAMPUSES[0]
C2 = CAMPUSES[1 % len(CAMPUSES)]
C3 = CAMPUSES[2 % len(CAMPUSES)]

HDR = ['活动编号', '活动名称', '活动级别', '主办单位', '开展时间', '地点', '负责人',
       '长号&短号', '活动类型', '项目类别', '活动对象', '参加方式', '备注', '校区', '加分分值']

# 5 份单表：每份都对应一个要演示的判定
SHEETS = [
    ('建工学院.xlsx', '第4周活动登记表', [
        # 正常活动（跨天）→ 演示「持续时间长的先排」
        ['', '2026年建筑工程学院大学生素质拓展中心结构设计大赛', '院级',
         '建筑工程学院大学生素质拓展中心', '周一09:00-周三00:00', '二教A201',
         '示例·甲', '13800000001', '临时', '学术科技与专业技能', '全院学生',
         '线下', '', C1, '0.2-1.0'],
    ], '演示：正常活动 + 跨天时长排序'),

    ('化工学院.xlsx', '第4周活动登记表', [
        # 与建工那份重复（名称+时间+地点全同）→ 演示「精确重复自动合并」
        ['', '2026年建筑工程学院大学生素质拓展中心结构设计大赛', '院级',
         '建筑工程学院大学生素质拓展中心', '周一09:00-周三00:00', '二教A201',
         '示例·甲', '13800000001', '临时', '学术科技与专业技能', '全院学生',
         '线下', '', C1, '0.2-1.0'],
        # 协会会员培训 → 演示「表一剔除协会会员」
        ['', '2026年校青年志愿者协会会员培训', '校级', '校青年志愿者协会',
         '周日14:30-17:30', '图书馆报告厅', '示例·乙', '13800000002', '常规',
         '社会实践与志愿服务', '协会会员', '线下', '', C2, '0.2'],
    ], '演示：跨来源重复上报 + 协会会员活动'),

    ('体育学院.xlsx', '第4周活动登记表', [
        # 跨周长活动 → 演示「15 天排最前」
        ['', '2026年体育学院篮球联赛', '院级', '体育学院学生会',
         '第三周周三-第五周周三', '风雨球场', '示例·丙', '13800000003', '临时',
         '文化艺术与体育竞技', '全院学生', '线下', '', C1, '0.5'],
    ], '演示：跨周活动（持续 15 天，排最前）'),

    ('外国语学院.xlsx', '第4周活动登记表', [
        # 故意写错：引号 / 非法项目类别 / 全角标点 / 级院对象不符 / 线上无链接
        ['', '2026年外国语学院"英语角"活动', '院（系）级', '外国语学院学生会',
         '周三19:30-20:30', '线上', '示例·丁', '13800000004', '临时', '学术科技',
         '全校学生', '线上', '备注：请提前报名，名额有限', C3, '0.2'],
    ], '演示：校验能挑出 5 类问题（引号/类别/全角标点/级别对象/线上无链接）'),

    ('理学院.xlsx', '第5周活动登记表', [
        # 属于第 5 周 → 演示「绝不跨周混表」
        ['', '2026年理学院数学建模竞赛', '院级', '理学院学生会',
         '周六08:00-12:00', '实训楼302', '示例·戊', '13800000005', '临时',
         '学术科技与专业技能', '全院学生', '线下', '', C2, '0.2-0.8'],
    ], '演示：混进来的第 5 周文件（按周分组，不混表）'),
]


def make_table(path: Path, title: str, rows):
    """写成和学校模板同样的形态：标题行 + 表头行 + 数据行，宋体、有边框。"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    thin = Side(style='thin', color='FF000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=15)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name='宋体', size=24, bold=True, charset=134)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    for i, h in enumerate(HDR, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(name='宋体', size=9, bold=True, charset=134)
        c.border = border
        c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 21

    for r, row in enumerate(rows, start=3):
        for i, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=i, value=(v if v != '' else None))
            cell.font = Font(name='宋体', size=9, charset=134)
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[r].height = 21

    for i in range(1, 16):
        ws.column_dimensions[chr(64 + i)].width = 18
    ws.column_dimensions['B'].width = 52
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def make_appdoc(path: Path):
    """故意留下一批必改项，用来演示预审工具。"""
    d = Document()
    d.add_paragraph('临时项目申报书')          # 故意缺「大学生素质拓展中心制」
    d.add_paragraph('项目名称：2026年"计算机学院"程序设计竞赛')   # 引号（必改）
    d.add_paragraph('实施时间：2026.3.10')                      # 非年月日（必改）
    d.add_paragraph('活动流程：报名后进行比赛，评出一二三等奖。')  # 字数不足（必改）
    # 故意不写「工作人员不可加分」与证件检查句（两条必改）
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = '一等奖'
    t.cell(0, 1).text = '30%'
    t.cell(1, 0).text = '二等奖'
    t.cell(1, 1).text = '30%'                                  # 合计 60%（必改）
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


def main():
    if DEMO.exists():
        shutil.rmtree(DEMO)
    INBOX.mkdir(parents=True)

    print('=== 生成单表 ===')
    for name, title, rows, why in SHEETS:
        p = INBOX / name
        make_table(p, title, rows)
        print(f'  {name:<20} {p.stat().st_size:>6} B   {why}')

    print('\n=== 生成待预审申报书 ===')
    make_appdoc(APPDOC)
    print(f'  {APPDOC.name:<20} {APPDOC.stat().st_size:>6} B   '
          f'故意含 6 处必改项')

    print('\n=== 自检：用真实内核跑一遍，确认演示效果 ===')
    sys.path.insert(0, str(HERE))
    from core import zjb, merge as mergemod, agent as agentmod

    cfg = zjb.load_config(HERE / 'core' / 'config.json')
    r = mergemod.run_merge([INBOX], DEMO / '_自检输出', cfg)
    c = r['counts']
    print(f'  周次 {r["week"]}｜单表 {c["input_files"]} 份｜原始行 {c["raw_rows"]}'
          f'｜去重后 {c["unique"]} 个活动')
    print(f'  自动合并重复 {c["exact_merged"]} 条｜表一 {c["publicity_rows"]} 行'
          f'｜校验错误 {c["activity_errors"]} 处 / 提醒 {c["activity_warnings"]} 处')
    print(f'  跨周保护：{len(r.get("extra_weeks") or [])} 套额外周次产物'
          f'｜提示 {len(r.get("notes") or [])} 条')
    for n in (r.get('notes') or []):
        print('    ·', n)
    print('  排序（演示「长的先排」）:')
    for a in r['activities']:
        print(f'    {a["no"]}. 持续{a["duration_days"]:>2}天  {a["name"][:30]}')

    pr = agentmod.run('预审这份申报书', doc=str(APPDOC))
    chk = pr['results']['pre_review_application']
    print(f'\n  预审自检：必改 {len(chk["must_fix"])} 项，无法判断 {len(chk["cannot_judge"])} 项')
    for m in chk['must_fix']:
        print('    ✗', m)

    shutil.rmtree(DEMO / '_自检输出', ignore_errors=True)
    print(f'\n演示数据就绪：{DEMO}')
    print('  收件箱 →', INBOX)
    print('  申报书 →', APPDOC)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
