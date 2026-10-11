#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
质检部周工作自动化 —— 核心脚本
广东石油化工学院 · 校素拓质检部（官渡校区）

子命令:
  info       查看登记表结构、活动清单、主办学院识别结果
  split      按「校区」把总表拆成多份（官渡 / 西城）
  arrange    安排二级学院检查活动（含跨周防重复）
  assistant  安排助理检查活动（按空闲时间表）
  check      校验表格是否满足部内规则
  ledger     查看 / 维护跨周排班台账
  slots      生成本周空闲时间表模板

示例:
  python zjb.py info "第四周活动登记表(总表).xlsx"
  python zjb.py split "第四周活动登记表(总表).xlsx"
  python zjb.py arrange "第四周活动登记表（官渡）.xlsx" --week 4 --recommend
  python zjb.py arrange "第四周活动登记表（官渡）.xlsx" --week 4 --colleges 理学院,外国语学院,体育学院,艺术学院
  python zjb.py assistant "第四周活动登记表（官渡）（已安排二级学院）.xlsx" --week 4 --slots slots.json
"""
from __future__ import annotations

import argparse
import json
import math
import os
import re
import shutil
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖 openpyxl，请先执行: pip install openpyxl")

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_CONFIG = SKILL_DIR / "assets" / "config.json"

# ---------------------------------------------------------------- 表头识别

HEADER_KEYS = {
    "no": ["活动编号", "编号", "序号"],
    "name": ["活动名称", "活动名", "项目名称"],
    "level": ["活动级别", "级别"],
    "host": ["主办单位", "主办方", "主办"],
    "time": ["开展时间", "活动时间"],
    "place": ["地点", "活动地点"],
    "contact": ["负责人"],
    "phone": ["长号&短号", "长号", "短号", "联系方式"],
    "type": ["活动类型"],
    "category": ["项目类别"],
    "target": ["活动对象", "对象"],
    "way": ["参加方式"],
    "remark": ["备注"],
    "campus": ["校区"],
    "score": ["加分分值", "分值", "加分"],
}

WEEKDAY_CN = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "日": 7, "天": 7}
CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9,
          "十": 10, "十一": 11, "十二": 12, "十三": 13, "十四": 14, "十五": 15, "十六": 16,
          "十七": 17, "十八": 18, "十九": 19, "二十": 20}
TRAILING_PAREN = re.compile(r"[（(][^（）()]{0,12}[）)]\s*$")


# ---------------------------------------------------------------- 基础工具

def load_config(path=None) -> dict:
    p = Path(path) if path else DEFAULT_CONFIG
    if not p.exists():
        raise FileNotFoundError(f"找不到配置文件: {p}")
    with open(p, encoding="utf-8") as f:
        cfg = json.load(f)
    cfg.setdefault("campuses", ["官渡", "西城"])
    cfg.setdefault("colleges", [])
    cfg.setdefault("campus_aliases", {})
    cfg.setdefault("college_campus", {})
    cfg.setdefault("checker_column", ["检查学院"])
    cfg.setdefault("assistant_column", ["检查助理"])
    cfg.setdefault("per_activity_colleges", 2)
    cfg.setdefault("per_activity_assistants", 2)
    cfg.setdefault("suffix_colleges", "（已安排二级学院）")
    cfg.setdefault("suffix_assistants", "（已安排助理）")
    cfg.setdefault("value_separator", "、")
    cfg.setdefault("last_week_scope", "global")
    cfg.setdefault("ledger_path", "data/ledger.json")
    cfg.setdefault("standardize_name", True)
    cfg.setdefault("term_name", "")
    cfg.setdefault("file_name_template", "{term}第{week}周素质拓展活动登记表")
    cfg.setdefault("campus_paren", "")
    cfg.setdefault("empty_week_note", "本周无活动")
    cfg["_path"] = str(p)
    return cfg


def ledger_file(cfg: dict) -> Path:
    p = Path(cfg["ledger_path"])
    return p if p.is_absolute() else SKILL_DIR / p


def load_ledger(cfg: dict) -> dict:
    p = ledger_file(cfg)
    if not p.exists():
        return {"version": 1, "weeks": {}}
    with open(p, encoding="utf-8") as f:
        data = json.load(f)
    data.setdefault("weeks", {})
    return data


def save_ledger(cfg: dict, data: dict) -> Path:
    p = ledger_file(cfg)
    p.parent.mkdir(parents=True, exist_ok=True)
    data["updated"] = datetime.now().isoformat(timespec="seconds")
    with open(p, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return p


def cn_num(text: str):
    text = str(text).strip()
    if text.isdigit():
        return int(text)
    return CN_NUM.get(text)


def extract_week(text: str):
    """从 '第4周活动登记表' / '第四周...' 里取出周次"""
    if not text:
        return None
    m = re.search(r"第\s*([0-9]{1,2})\s*周", str(text))
    if m:
        return int(m.group(1))
    m = re.search(r"第\s*([一二三四五六七八九十]{1,3})\s*周", str(text))
    if m:
        return cn_num(m.group(1))
    return None


def normalize_campus(value: str, cfg: dict):
    if value is None:
        return None
    v = str(value).strip()
    if not v:
        return None
    return cfg["campus_aliases"].get(v, v)


def find_header_row(ws):
    """定位表头行，返回 (row_idx, {key: col_idx})"""
    limit = min(ws.max_row, 8)
    for r in range(1, limit + 1):
        mapping = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            s = str(v).strip().replace(" ", "")
            for key, aliases in HEADER_KEYS.items():
                if key in mapping:
                    continue
                for a in aliases:
                    if s == a.replace(" ", "") or (a in s and len(s) <= len(a) + 4):
                        mapping[key] = c
                        break
        if "name" in mapping and ("campus" in mapping or "host" in mapping):
            return r, mapping
    raise ValueError("未能识别表头行（需包含「活动名称」以及「校区」或「主办单位」列）")


def act_key(a: dict) -> str:
    """活动的稳定主键：优先活动编号，其次名称，再补时间地点区分同名活动。

    借鉴 OfficeCLI 的「稳定 ID 寻址」：不依赖行号。表里插入/删除行、标题行增减
    都不会让台账与表格对不上号。
    """
    base = (a.get("no") or "").strip() or (a.get("name") or "").strip()
    tail = "|".join(str(x).strip() for x in (a.get("time"), a.get("place")) if x)
    raw = base + ("|" + tail if tail else "")
    return re.sub(r"\s+", "", raw) or "未命名活动"


def read_activities(ws, header_row: int, header_map: dict, cfg: dict):
    acts = []
    seen = set()
    for r in range(header_row + 1, ws.max_row + 1):
        def cell(key):
            col = header_map.get(key)
            if not col:
                return None
            v = ws.cell(row=r, column=col).value
            return None if v is None else str(v).strip()

        name = cell("name")
        no = cell("no")
        if not name and not no:
            continue
        host = cell("host") or ""
        campus = normalize_campus(cell("campus"), cfg)
        item = {
            "row": r,
            "no": no,
            "name": name or "",
            "host": host,
            "host_college": match_college(host, cfg["colleges"]),
            "time": cell("time") or "",
            "place": cell("place") or "",
            "campus": campus,
            "level": cell("level") or "",
        }
        k = act_key(item)
        uniq, n = k, 2
        while uniq in seen:          # 同表内重名活动时补序号，保证主键唯一
            uniq = f"{k}#{n}"
            n += 1
        seen.add(uniq)
        item["key"] = uniq
        acts.append(item)
    return acts


def match_college(text: str, colleges):
    """从「化学工程学院大学生素质拓展中心」中匹配出「化学工程学院」（最长优先）"""
    if not text:
        return None
    for c in sorted(colleges, key=len, reverse=True):
        if c in str(text):
            return c
    return None


def parse_slots(text: str):
    """把「周日08:00-12:00 14:30-17:30」解析成 ['周日上午','周日下午']"""
    if not text:
        return []
    t = str(text)
    day = None
    if re.search(r"周[六日天]|星期[六日天]", t):
        m = re.search(r"周([六日天])|星期([六日天])", t)
        ch = m.group(1) or m.group(2)
        day = "周六" if ch == "六" else "周日"
    if day is None:
        return []
    hours = [int(h) for h, _ in re.findall(r"(\d{1,2})[:：](\d{2})", t)]
    starts = hours[0::2]
    slots = []
    for h in starts:
        if h <= 12:
            s = f"{day}上午"
        elif h < 18:
            s = f"{day}下午"
        else:
            continue
        if s not in slots:
            slots.append(s)
    if not slots:
        if "上午" in t or "早上" in t:
            slots.append(f"{day}上午")
        if "下午" in t:
            slots.append(f"{day}下午")
    return slots


def last_used_col(ws, header_row: int, header_map: dict) -> int:
    """表格实际占用的最后一列（表头行 + 数据区都要看）。

    官方格式里第二个检查学院列没有表头（P 有表头、Q 只填数据），
    只扫表头行会把它当成空列，导致后续新增列覆盖已有数据。
    """
    last = max(header_map.values()) if header_map else 1
    for c in range(1, ws.max_column + 1):
        hv = ws.cell(row=header_row, column=c).value
        if hv is not None and str(hv).strip():
            last = max(last, c)
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip():
                last = max(last, c)
                break
    return last


def ensure_column(ws, header_row: int, header_map: dict, aliases, title: str, cfg: dict):
    """找到已存在的目标列；没有则在表格末尾新增一列（复制表头样式）"""
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v and str(v).strip() in aliases:
            return c, False
    # 新列位置基于表头行 + 数据区实际占用的最大列，否则会覆盖已有的无表头数据列
    last = last_used_col(ws, header_row, header_map)
    new_col = last + 1
    src = ws.cell(row=header_row, column=last)
    tgt = ws.cell(row=header_row, column=new_col, value=title)
    tgt._style = copy(src._style)
    ws.column_dimensions[get_column_letter(new_col)].width = max(14, len(title) * 3.2)
    return new_col, True


def read_checkers(ws, header_row, aliases, cfg):
    """读取检查方数据列，兼容两种格式：
    ① 官方横向展开：表头仅起始列有字，数据向右连续多格、一格一学院
    ② 单格合并：一格内「学院A、学院B」
    返回 (col_start, {row: [学院, ...]})；未找到列返回 (None, {})
    """
    col_start = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v and str(v).strip() in aliases:
            col_start = c
            break
    if col_start is None:
        return None, {}
    result = {}
    for r in range(header_row + 1, ws.max_row + 1):
        vals = []
        for c in range(col_start, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None or not str(v).strip():
                break
            vals.append(str(v).strip())
        if vals:
            names = []
            for v in vals:
                names.extend(x.strip() for x in re.split(r"[,，、\s]+", v) if x.strip())
            result[r] = names
    return col_start, result


def ensure_checker_columns(ws, header_row, header_map, per, cfg):
    """安排写入用：找到已有检查列；没有则按官方格式新增横向展开的 per 列（一格一学院）。
    新列起点取「表头行与数据区实际占用的最大列 + 1」，避免覆盖已有数据。"""
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v and str(v).strip() in cfg["checker_column"]:
            return c, False
    last = last_used_col(ws, header_row, header_map)
    col_start = last + 1
    src = ws.cell(row=header_row, column=last)
    tgt = ws.cell(row=header_row, column=col_start, value=cfg["checker_column"][0])
    tgt._style = copy(src._style)
    ws.column_dimensions[get_column_letter(col_start)].width = max(
        14, len(cfg["checker_column"][0]) * 2.2)
    for i in range(1, per):
        ws.column_dimensions[get_column_letter(col_start + i)].width = 14
    return col_start, True


def extend_title_merge(ws, last_col: int):
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row == 1 and rng.min_col == 1:
            ws.unmerge_cells(str(rng))
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(last_col, rng.max_col))
            return True
    return False


def write_values(ws, col: int, values_by_row: dict, cfg: dict):
    sep = cfg["value_separator"]
    for row_idx, values in values_by_row.items():
        cell = ws.cell(row=row_idx, column=col)
        cell.value = sep.join(values) if isinstance(values, (list, tuple)) else values
        cell._style = copy(ws.cell(row=row_idx, column=1)._style)


STATUS_SUFFIXES = {
    "（已安排二级学院）": ["（已安排二级学院）", "(已安排二级学院)"],
    "（已安排助理）": ["（已安排助理）", "(已安排助理)"],
}


def _canon_status(s: str) -> str:
    for canon, variants in STATUS_SUFFIXES.items():
        if s in variants:
            return canon
    return s


PAREN_WRAP = re.compile(r"^[（(]\s*(.+?)\s*[）)]$")
TAIL_PAREN = re.compile(r"([（(])\s*([^（）()]{1,12}?)\s*([）)])\s*$")
STATUS_ORDER = list(STATUS_SUFFIXES.keys())


def unwrap_paren(s: str):
    """「（官渡）」→「官渡」；不是括号包裹则返回 None"""
    m = PAREN_WRAP.match(str(s).strip())
    return m.group(1).strip() if m else None


def is_campus_name(name, cfg):
    if not name or not cfg:
        return False
    campuses = cfg.get("campuses") or []
    aliases = cfg.get("campus_aliases") or {}
    return name in campuses or aliases.get(name, name) in campuses


def out_path_for(src: Path, add_suffix=None, base_name=None, outdir=None, cfg=None, **_ignored):
    """生成输出文件名（校区后缀与状态后缀分开处理，避免同名互相覆盖）

    · 状态后缀（已安排二级学院 / 已安排助理）按固定顺序累积在原位
    · 校区后缀排在状态后缀之前；括号风格自适应源文件名（半角 (西城) / 全角（西城））
    · 尾部非校区括号（如「(总表)」）会被剥掉
    · 若输出路径与源文件相同，自动追加「（副本）」防止覆盖

    例：第四周活动登记表（官渡）→（官渡）（已安排二级学院）→（官渡）（已安排二级学院）（已安排助理）
    """
    stem = src.stem

    # 1) 反复剥掉尾部的状态后缀
    present = []
    while True:
        s = stem.rstrip()
        hit = None
        for canon, variants in STATUS_SUFFIXES.items():
            for v in variants:
                if s.endswith(v):
                    hit = (canon, len(v))
                    break
            if hit:
                break
        if not hit:
            break
        stem = s[: -hit[1]].rstrip()
        if hit[0] not in present:
            present.append(hit[0])

    # 2) 处理尾部括号：是校区就记住名字+括号风格，不是校区（如「总表」）直接剥掉
    campus_old, style = None, ("（", "）")
    m = TAIL_PAREN.search(stem)
    if m:
        inner = m.group(2).strip()
        if is_campus_name(inner, cfg):
            campus_old = inner
            style = (m.group(1), m.group(3))
        stem = stem[: m.start()].rstrip()

    # 2.1) 括号风格可按配置强制统一（用户习惯：半角 "(西城)"）
    forced = ((cfg or {}).get("campus_paren") or "").strip()
    if forced:
        style = (forced, ")" if forced == "(" else "）")

    # 3) 解析本次追加的后缀：校区 还是 状态
    campus_new, add_status = None, None
    if add_suffix:
        inner = unwrap_paren(add_suffix)
        if is_campus_name(inner, cfg):
            campus_new = inner
        else:
            canon = _canon_status(add_suffix)
            add_status = canon if canon in STATUS_SUFFIXES else None

    campus = campus_new or campus_old
    if add_status and add_status not in present:
        present.append(add_status)
    if base_name:
        stem = base_name

    name = stem
    if campus:
        name += f"{style[0]}{campus}{style[1]}"
    name += "".join(s for s in STATUS_ORDER if s in present)

    d = Path(outdir) if outdir else src.parent
    out = d / f"{name}.xlsx"
    if out.resolve() == src.resolve():
        out = d / f"{name}（副本）.xlsx"
    return out


def jdump(obj, args):
    if getattr(args, "json", False):
        print(json.dumps(obj, ensure_ascii=False, indent=2))
    else:
        return False
    return True


def hr(title):
    print("\n" + "=" * 60)
    print(title)
    print("=" * 60)


# ---------------------------------------------------------------- info

def cmd_info(args):
    cfg = load_config(args.config)
    src = Path(args.xlsx)
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = find_header_row(ws)
    acts = read_activities(ws, header_row, header_map, cfg)
    week = extract_week(ws.cell(row=1, column=1).value) or extract_week(src.name)
    result = {
        "file": str(src),
        "sheet": ws.title,
        "title": ws.cell(row=1, column=1).value,
        "week": week,
        "header_row": header_row,
        "header_map": {k: get_column_letter(v) for k, v in header_map.items()},
        "columns": {k: ws.cell(row=header_row, column=v).value for k, v in header_map.items()},
        "activity_count": len(acts),
        "campuses": sorted({a["campus"] for a in acts if a["campus"]}),
        "activities": acts,
        "host_colleges": sorted({a["host_college"] for a in acts if a["host_college"]}),
    }
    if jdump(result, args):
        return
    hr(f"登记表结构 · {src.name}")
    print(f"标题      : {result['title']}")
    print(f"周次      : 第{week}周" if week else "周次      : 未识别")
    print(f"表头行    : 第 {header_row} 行")
    print(f"列映射    : " + ", ".join(f"{k}→{get_column_letter(v)}" for k, v in header_map.items()))
    print(f"活动数量  : {len(acts)}")
    print(f"涉及校区  : {'、'.join(result['campuses']) or '未填写'}")
    print(f"主办学院  : {'、'.join(result['host_colleges']) or '未识别'}")
    print("\n活动清单：")
    for a in acts:
        print(f"  [{a['no']}] {a['name']}")
        print(f"       主办={a['host_college'] or a['host'] or '-'} | 时间={a['time'] or '-'} | "
              f"地点={a['place'] or '-'} | 校区={a['campus'] or '-'}")


# ---------------------------------------------------------------- split

def clone_subset(src_ws, keep_rows, last_col=None):
    """把 src_ws 中指定行（按顺序重排）克隆到新工作表，保留样式/行高/列宽/合并"""
    import openpyxl as _oxl
    wb = _oxl.Workbook()
    ws = wb.active
    ws.title = src_ws.title
    last_col = last_col or src_ws.max_column
    for key, dim in src_ws.column_dimensions.items():
        if dim.width or dim.hidden:
            nd = ws.column_dimensions[key]
            nd.width = dim.width
            nd.hidden = dim.hidden
    row_map = {old: new for new, old in enumerate(keep_rows, start=1)}
    for old, new in row_map.items():
        src_dim = src_ws.row_dimensions.get(old)
        if src_dim is not None and src_dim.height:
            ws.row_dimensions[new].height = src_dim.height
        for c in range(1, last_col + 1):
            sc = src_ws.cell(row=old, column=c)
            tc = ws.cell(row=new, column=c)
            tc.value = sc.value
            if sc.has_style:
                tc.font = copy(sc.font)
                tc.fill = copy(sc.fill)
                tc.border = copy(sc.border)
                tc.alignment = copy(sc.alignment)
                tc.number_format = sc.number_format
                tc.protection = copy(sc.protection)
    for rng in src_ws.merged_cells.ranges:
        if rng.min_row in row_map and rng.max_row in row_map:
            ws.merge_cells(start_row=row_map[rng.min_row], start_column=rng.min_col,
                           end_row=row_map[rng.max_row], end_column=min(rng.max_col, last_col))
    return wb, ws


def campus_in_name(src: Path, cfg: dict):
    """从文件名尾部括号识别校区（先剥掉状态后缀，避免误判）"""
    s = src.stem
    while True:
        changed = False
        for _canon, variants in STATUS_SUFFIXES.items():
            for v in variants:
                if s.rstrip().endswith(v):
                    s = s.rstrip()[: -len(v)]
                    changed = True
        if not changed:
            break
    m = TAIL_PAREN.search(s)
    if m and is_campus_name(m.group(2).strip(), cfg):
        return m.group(2).strip()
    return None


def standard_base_name(cfg: dict, week):
    """按配置生成标准文件名主干：2026-2027学年第一学期第5周素质拓展活动登记表"""
    if not cfg.get("standardize_name", True) or week is None:
        return None
    tpl = cfg.get("file_name_template") or ""
    if not tpl:
        return None
    try:
        return (tpl.format(term=cfg.get("term_name") or "", week=week)).strip() or None
    except (KeyError, IndexError, ValueError):
        return None


def fill_empty_week(ws, src_ws, header_row: int, last_col: int, header_map: dict, cfg: dict):
    """某校区本周无活动：补一行沿用原表样式的空行，并在备注列写明"""
    tpl_row = header_row + 1
    if tpl_row > src_ws.max_row:
        return False
    new_row = header_row + 1
    for c in range(1, last_col + 1):
        sc = src_ws.cell(row=tpl_row, column=c)
        tc = ws.cell(row=new_row, column=c)
        if sc.has_style:
            tc.font = copy(sc.font)
            tc.fill = copy(sc.fill)
            tc.border = copy(sc.border)
            tc.alignment = copy(sc.alignment)
            tc.number_format = sc.number_format
            tc.protection = copy(sc.protection)
    dim = src_ws.row_dimensions.get(tpl_row)
    if dim is not None and dim.height:
        ws.row_dimensions[new_row].height = dim.height
    col = header_map.get("remark") or header_map.get("name") or 2
    ws.cell(row=new_row, column=col).value = cfg.get("empty_week_note") or "本周无活动"
    return True


def cmd_split(args):
    cfg = load_config(args.config)
    src = Path(args.xlsx)
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = find_header_row(ws)
    campus_col = header_map.get("campus")
    if not campus_col:
        raise ValueError("表中没有「校区」列，无法拆分")
    acts = read_activities(ws, header_row, header_map, cfg)
    if not acts:
        raise ValueError("表里没有活动数据行，无法拆分")
    week = extract_week(ws.cell(row=1, column=1).value) or extract_week(src.name)

    groups = {}
    for a in acts:
        groups.setdefault(a["campus"] or "未知校区", []).append(a["row"])

    # 源表若已按校区拆过（文件名带校区），默认只处理该校区；否则按全校区清单出表
    known = campus_in_name(src, cfg)
    force_all = getattr(args, "all", False)
    if known and not force_all:
        targets = [known]
    else:
        targets = list(dict.fromkeys((cfg.get("campuses") or []) + list(groups.keys())))

    if args.outdir:
        Path(args.outdir).mkdir(parents=True, exist_ok=True)
    head_rows = list(range(1, header_row + 1))
    last_col = max(header_map.values())
    base = args.base_name or standard_base_name(cfg, week)

    outputs = []
    for campus in targets:
        rows = sorted(groups.get(campus, []))
        wb2, ws2 = clone_subset(ws, head_rows + rows, last_col=last_col)
        empty = not rows
        if empty:
            fill_empty_week(ws2, ws, header_row, last_col, header_map, cfg)
        out = out_path_for(src, f"（{campus}）", base_name=base, outdir=args.outdir, cfg=cfg)
        atomic_save(wb2, out)
        outputs.append({"campus": campus, "rows": len(rows), "empty": empty, "file": str(out)})

    if jdump({"source": str(src), "week": week, "outputs": outputs}, args):
        return
    hr(f"拆分完成 · {src.name}")
    if known and not force_all:
        print(f"  提示：源表已按校区拆分（{known}）；如需补齐其他校区的空白表，加 --all")
    for o in outputs:
        tag = (cfg.get("empty_week_note") or "本周无活动") if o["empty"] else f"{o['rows']} 个活动"
        print(f"  {o['campus']:<6} {tag} → {o['file']}")


# ---------------------------------------------------------------- arrange

def last_week_checkers(ledger: dict, week: int, cfg: dict):
    """取第 week-1 周已被安排过的学院（默认跨校区合并）"""
    if week is None:
        return set(), []
    prev = week - 1
    if prev < 1:
        return set(), []
    scope = cfg["last_week_scope"]
    found = set()
    detail = []
    for key, rec in ledger.get("weeks", {}).items():
        try:
            w = int(rec.get("week"))
        except (TypeError, ValueError):
            continue
        if w != prev:
            continue
        if scope == "campus" and rec.get("campus") != cfg.get("_current_campus"):
            continue
        for name in rec.get("checkers", []):
            found.add(name)
        for act in rec.get("activities", []):
            for name in act.get("checkers", []):
                found.add(name)
        detail.append(f"第{w}周 · {rec.get('campus', '-')}")
    return found, detail


def record_to_ledger(ledger, cfg, week, campus, acts, plan, kind="colleges"):
    """把本周安排写进台账（同一周同一校区重复执行时按行合并）"""
    key = f"W{week}-{campus}" if week else f"W?-{campus}"
    weeks = ledger.setdefault("weeks", {})
    entry = weeks.setdefault(key, {
        "week": week, "campus": campus,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "activities": [],
    })
    entry["updated_at"] = datetime.now().isoformat(timespec="seconds")
    by_key = {x.get("key"): x for x in entry["activities"] if x.get("key")}
    # 按表格行顺序遍历，保证台账顺序稳定且与表一致（不再按行号排序）
    for a in acts:
        row = a["row"]
        names = plan.get(row)
        if names is None:
            continue
        k = a.get("key") or f"row{row}"
        item = by_key.get(k)
        if item is None:
            # 兼容旧台账（只有 row / no+name，没有 key）
            item = next((x for x in entry["activities"]
                         if (x.get("no") or "") == (a.get("no") or "")
                         and (x.get("name") or "") == (a.get("name") or "")), None)
        if item is None:
            item = {}
            entry["activities"].append(item)
        by_key[k] = item
        item.update({
            "key": k, "row": row,
            "no": a.get("no"), "name": a.get("name"), "host_college": a.get("host_college"),
            "time": a.get("time"), "campus": campus,
        })
        if kind == "colleges":
            item["checkers"] = names
        else:
            item["assistants"] = names
    all_checkers = set(entry.get("checkers") or [])
    for item in entry["activities"]:
        all_checkers.update(item.get("checkers") or [])
    entry["checkers"] = sorted(all_checkers)
    return key


def backup_file(path: Path):
    """写入前把目标文件备份为 xxx.bak.xlsx（只保留最近一份）"""
    try:
        if path.exists():
            bak = path.with_name(path.stem + ".bak" + path.suffix)
            shutil.copy2(path, bak)
            return bak
    except Exception:
        pass
    return None


def atomic_save(wb, out: Path):
    """原子落盘：先写临时文件，成功后再 os.replace 覆盖。

    借鉴 OfficeCLI batch 的原子性：save 中途出错（磁盘满、文件被占用）时，
    原文件分毫不动，不会留下打不开的半截表。落盘前自动备份一份 .bak。
    """
    out.parent.mkdir(parents=True, exist_ok=True)
    bak = backup_file(out)
    tmp = out.with_name(out.name + ".tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, out)
    except Exception:
        if tmp.exists():
            try:
                tmp.unlink()
            except Exception:
                pass
        raise
    return bak


def recommend_pool(acts, cfg, ledger, week):
    last, last_detail = last_week_checkers(ledger, week, cfg)
    hosts = {a["host_college"] for a in acts if a["host_college"]}
    cands = []
    for c in cfg["colleges"]:
        cands.append({
            "college": c,
            "last_week": c in last,
            "host_this_week": c in hosts,
        })
    need = max(4, math.ceil(len(acts) * cfg["per_activity_colleges"] / 2)) if acts else 4
    fresh = [c["college"] for c in cands if not c["last_week"]]
    return {
        "week": week,
        "last_week_checkers": sorted(last),
        "last_week_records": last_detail,
        "host_colleges": sorted(hosts),
        "candidates": cands,
        "available_now": fresh,
        "suggested_pool_size": need,
        "rule_note": f"每活动 {cfg['per_activity_colleges']} 个学院；举办学院不检查本学院；同一学院不连续两周",
    }


def cmd_arrange(args):
    cfg = load_config(args.config)
    src = Path(args.xlsx)
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = find_header_row(ws)
    acts = read_activities(ws, header_row, header_map, cfg)
    if not acts:
        raise ValueError("表中没有读到任何活动行，请检查表格内容")
    week = args.week or extract_week(ws.cell(row=1, column=1).value) or extract_week(src.name)
    campuses = {a["campus"] for a in acts if a["campus"]}
    if len(campuses) > 1:
        print(f"⚠️ 该表包含多个校区（{'、'.join(sorted(campuses))}），建议先执行 split 拆分后再安排；"
              f"本次将按第一个校区（{acts[0].get('campus')}）记录台账。", file=sys.stderr)
    campus = normalize_campus(acts[0].get("campus"), cfg) or "未知"
    cfg["_current_campus"] = campus
    ledger = load_ledger(cfg)

    if args.recommend:
        info = recommend_pool(acts, cfg, ledger, week)
        if jdump(info, args):
            return
        hr(f"安排候选推荐 · 第{week}周 · {campus}校区")
        print(f"本周主办学院（不能检查自己）: {'、'.join(info['host_colleges']) or '-'}")
        print(f"上周已安排（不连续两周）    : {'、'.join(info['last_week_checkers']) or '无记录'}")
        print(f"\n可选学院候选（★=上周没用过）：")
        for c in info["candidates"]:
            mark = "  " if c["last_week"] else "★ "
            tag = " [本周主办]" if c["host_this_week"] else ""
            print(f"  {mark}{c['college']}{tag}")
        print(f"\n建议安排学院数量: {info['suggested_pool_size']} 个（{len(acts)} 个活动 × "
              f"{cfg['per_activity_colleges']} 个学院）")
        return

    pool = []
    if args.colleges:
        pool = [c.strip() for c in re.split(r"[,，、\s]+", args.colleges) if c.strip()]
    if not pool:
        raise SystemExit("请用 --colleges 指定本周安排的学院，或先用 --recommend 查看候选")
    unknown = [c for c in pool if c not in cfg["colleges"]]
    if unknown:
        print(f"⚠️ 以下学院不在名单内（仍继续，建议核对 config.json）：{'、'.join(unknown)}", file=sys.stderr)

    last, _ = last_week_checkers(ledger, week, cfg)
    per = cfg["per_activity_colleges"]
    load = {c: 0 for c in pool}
    plan = {}
    problems = []
    for a in acts:
        host = a["host_college"]
        cand = [c for c in pool if c != host]
        if len(cand) < per:
            problems.append(f"活动「{a['name']}」可选学院不足 {per} 个（主办={host}，池={pool}）")
            continue
        cand.sort(key=lambda c: (1 if c in last else 0, load[c], c))
        chosen = cand[:per]
        for c in chosen:
            load[c] += 1
        plan[a["row"]] = chosen
        if any(c in last for c in chosen):
            problems.append(f"活动「{a['name']}」安排了上周已检查的学院："
                            + "、".join(c for c in chosen if c in last))
    if problems and not args.force:
        print("⚠️ 发现规则问题，已停止写入（确认无误可加 --force 强制写入）：", file=sys.stderr)
        for p in problems:
            print("   - " + p, file=sys.stderr)
        alt = [c for c in cfg["colleges"] if c not in last and c not in pool]
        if alt and any("上周已检查" in p for p in problems):
            print("   建议：把本周学院池换/扩成下列「上周未安排过」的学院 ——", file=sys.stderr)
            print("         " + "、".join(alt), file=sys.stderr)
        raise SystemExit(1)

    col_idx, is_new = ensure_checker_columns(ws, header_row, header_map, per, cfg)
    # 官方格式：横向展开，一格一学院（P、Q…各填一个）
    for row_idx, names in plan.items():
        for i, name in enumerate(names):
            cell = ws.cell(row=row_idx, column=col_idx + i)
            cell.value = name
            cell._style = copy(ws.cell(row=row_idx, column=1)._style)
    if is_new:
        extend_title_merge(ws, col_idx + per - 1)

    out = out_path_for(src, cfg["suffix_colleges"], base_name=args.base_name, outdir=args.outdir, cfg=cfg)
    if not args.dry_run:
        atomic_save(wb, out)
        record_to_ledger(ledger, cfg, week, campus, acts, plan, "colleges")
        save_ledger(cfg, ledger)

    result = {
        "file": str(out if not args.dry_run else src),
        "written": not args.dry_run,
        "week": week, "campus": campus,
        "column": get_column_letter(col_idx),
        "column_is_new": is_new,
        "plan": [{"row": a["row"], "no": a["no"], "name": a["name"],
                  "host_college": a["host_college"], "checkers": plan.get(a["row"], [])} for a in acts],
        "pool": pool,
        "load": load,
        "last_week_checkers": sorted(last),
        "problems": problems,
    }
    if jdump(result, args):
        return
    hr(f"二级学院安排完成 · 第{week}周 · {campus}校区")
    for item in result["plan"]:
        print(f"  [{item['no']}] {item['name']}")
        print(f"       主办 {item['host_college'] or '-'}  →  检查 {'、'.join(item['checkers']) or '未安排'}")
    print(f"\n各学院承担量: " + "、".join(f"{k}{v}" for k, v in load.items()))
    if problems:
        print("⚠️ 提示: " + "；".join(problems))
    print(f"\n产出文件: {out}" + ("" if not args.dry_run else "（dry-run，未写入）"))
    print(f"台账已记录: {ledger_file(cfg)}")


# ---------------------------------------------------------------- assistant

def cmd_assistant(args):
    cfg = load_config(args.config)
    src = Path(args.xlsx)
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = find_header_row(ws)
    acts = read_activities(ws, header_row, header_map, cfg)
    week = args.week or extract_week(ws.cell(row=1, column=1).value) or extract_week(src.name)
    campus = normalize_campus(acts[0].get("campus"), cfg) or "未知"
    cfg["_current_campus"] = campus

    if args.slots:
        slots = json.loads(Path(args.slots).read_text(encoding="utf-8"))
    elif args.slots_json:
        slots = json.loads(args.slots_json)
    else:
        raise SystemExit("请用 --slots <空闲时间表.json> 或 --slots-json '{...}' 提供空闲时间")

    per = cfg["per_activity_assistants"]
    load = {}
    plan = {}
    problems = []
    for a in acts:
        need = parse_slots(a["time"])
        if not need:
            problems.append(f"活动「{a['name']}」时间「{a['time']}」不是周末时段，跳过")
            continue
        sets = [set(slots.get(s, [])) for s in need]
        avail = set.intersection(*sets) if sets else set()
        loose = False
        if len(avail) < per:
            union = set.union(*sets) if sets else set()
            if len(union) >= per:
                avail, loose = union, True
        if len(avail) < per:
            problems.append(f"活动「{a['name']}」可用助理不足（需{per}人，时段={'+'.join(need)}，可用={sorted(avail)}）")
            avail_sorted = sorted(avail, key=lambda p: (load.get(p, 0), p))
        else:
            avail_sorted = sorted(avail, key=lambda p: (load.get(p, 0), p))
        chosen = avail_sorted[:per]
        for p in chosen:
            load[p] = load.get(p, 0) + 1
        plan[a["row"]] = chosen
        if loose and chosen:
            problems.append(f"活动「{a['name']}」按宽松匹配安排（无人全程有空）")

    col_idx, is_new = ensure_column(ws, header_row, header_map,
                                    cfg["assistant_column"], cfg["assistant_column"][0], cfg)
    write_values(ws, col_idx, plan, cfg)
    if is_new:
        extend_title_merge(ws, col_idx)
    out = out_path_for(src, cfg["suffix_assistants"], base_name=args.base_name, outdir=args.outdir, cfg=cfg)
    if not args.dry_run:
        atomic_save(wb, out)
    result = {
        "file": str(out if not args.dry_run else src),
        "written": not args.dry_run,
        "week": week, "campus": campus,
        "column": get_column_letter(col_idx),
        "plan": [{"row": a["row"], "no": a["no"], "name": a["name"], "slots": parse_slots(a["time"]),
                  "assistants": plan.get(a["row"], [])} for a in acts],
        "load": load, "problems": problems,
    }
    if jdump(result, args):
        return
    hr(f"助理安排完成 · 第{week}周 · {campus}校区")
    for item in result["plan"]:
        print(f"  [{item['no']}] {item['name']}  ({'+'.join(item['slots']) or '无周末时段'})")
        print(f"       检查助理: {'、'.join(item['assistants']) or '未安排'}")
    print(f"\n个人承担量: " + ("、".join(f"{k}{v}" for k, v in load.items()) or "无"))
    if problems:
        print("⚠️ 提示: " + "；".join(problems))
    print(f"\n产出文件: {out}" + ("" if not args.dry_run else "（dry-run，未写入）"))


# ---------------------------------------------------------------- check

def cmd_check(args):
    cfg = load_config(args.config)
    src = Path(args.xlsx)
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = find_header_row(ws)
    acts = read_activities(ws, header_row, header_map, cfg)
    week = args.week or extract_week(ws.cell(row=1, column=1).value) or extract_week(src.name)
    campus = normalize_campus(acts[0].get("campus"), cfg) or "未知"
    cfg["_current_campus"] = campus
    col, checker_map = read_checkers(ws, header_row, cfg["checker_column"], cfg)
    ledger = load_ledger(cfg)
    last, _ = last_week_checkers(ledger, week, cfg)
    issues = []
    for a in acts:
        names = checker_map.get(a["row"], [])
        if len(names) != cfg["per_activity_colleges"]:
            issues.append({"row": a["row"], "name": a["name"], "type": "数量不符",
                           "detail": f"应为{cfg['per_activity_colleges']}个学院，实际{len(names)}个"})
        if a["host_college"] and a["host_college"] in names:
            issues.append({"row": a["row"], "name": a["name"], "type": "自检",
                           "detail": f"主办学院 {a['host_college']} 不能检查自己"})
        dup = [n for n in names if n in last]
        if dup:
            issues.append({"row": a["row"], "name": a["name"], "type": "连续两周",
                           "detail": f"上周已安排：{'、'.join(dup)}"})
    result = {"file": str(src), "week": week, "campus": campus,
              "checker_column": get_column_letter(col) if col else None,
              "ok": not issues, "issues": issues}
    if jdump(result, args):
        return
    hr(f"规则校验 · {src.name}")
    if not col:
        print("未找到检查学院列，跳过校验")
    if not issues:
        print("✅ 全部通过：每活动 2 个学院、无自检、无连续两周")
    else:
        for i in issues:
            print(f"  ⚠️ 第{i['row']}行 [{i['type']}] {i['name']} — {i['detail']}")


# ---------------------------------------------------------------- ingest

def cmd_ingest(args):
    """把用户已排好检查学院/助理的表读入台账（只读表，不修改、不另存文件）"""
    cfg = load_config(args.config)
    src = Path(args.xlsx)
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = find_header_row(ws)
    acts = read_activities(ws, header_row, header_map, cfg)
    week = args.week or extract_week(ws.cell(row=1, column=1).value) or extract_week(src.name)
    campus = normalize_campus(acts[0].get("campus"), cfg) if acts else None
    if not campus:
        m = re.search(r"[（(](官渡|西城)[）)]", src.name)
        campus = m.group(1) if m else args.campus
    if not campus:
        raise SystemExit("无法确定校区（表内无活动、文件名无校区），请用 --campus 指定")
    cfg["_current_campus"] = campus
    c_col, checker_map = read_checkers(ws, header_row, cfg["checker_column"], cfg)
    a_col, asst_map = read_checkers(ws, header_row, cfg["assistant_column"], cfg)
    if not checker_map and not asst_map:
        raise SystemExit("表中没有找到检查学院/检查助理数据，请核对表头与数据列")
    ledger = load_ledger(cfg)
    if checker_map:
        record_to_ledger(ledger, cfg, week, campus, acts, checker_map, "colleges")
    if asst_map:
        record_to_ledger(ledger, cfg, week, campus, acts, asst_map, "assistants")
    save_ledger(cfg, ledger)
    if jdump({"file": str(src), "week": week, "campus": campus, "ingested": True}, args):
        return
    hr(f"台账入档 · 第{week}周 · {campus}校区 · {src.name}")
    for a in acts:
        print(f"  [{a['no']}] {a['name']}")
        if checker_map:
            print(f"       检查学院: {'、'.join(checker_map.get(a['row']) or []) or '未安排'}")
        if asst_map:
            print(f"       检查助理: {'、'.join(asst_map.get(a['row']) or []) or '未安排'}")
    print(f"\n台账已更新: {ledger_file(cfg)}")


# ---------------------------------------------------------------- ledger

def cmd_ledger(args):
    cfg = load_config(args.config)
    ledger = load_ledger(cfg)
    if args.reset:
        weeks = ledger.get("weeks", {})
        if args.week and args.campus:
            key = f"W{args.week}-{args.campus}"
            weeks.pop(key, None)
        elif args.week:
            for k in [k for k in weeks if k.startswith(f"W{args.week}-")]:
                weeks.pop(k, None)
        else:
            weeks.clear()
        save_ledger(cfg, ledger)
        print(f"台账已清理：{ledger_file(cfg)}")
        return
    data = {"ledger_file": str(ledger_file(cfg)), "weeks": ledger.get("weeks", {})}
    if jdump(data, args):
        return
    hr("质检部周工作台账")
    weeks = ledger.get("weeks", {})
    if not weeks:
        print("（暂无记录）")
        return
    for key in sorted(weeks, key=lambda k: (weeks[k].get("week") or 0, k)):
        rec = weeks[key]
        print(f"\n第{rec.get('week')}周 · {rec.get('campus')}  (记录于 {rec.get('updated_at') or rec.get('created_at')})")
        for a in rec.get("activities", []):
            print(f"   [{a.get('no')}] {a.get('name')}")
            print(f"        主办={a.get('host_college') or '-'}  检查学院={'、'.join(a.get('checkers') or []) or '-'}"
                  + (f"  助理={'、'.join(a.get('assistants') or [])}" if a.get('assistants') else ""))


# ---------------------------------------------------------------- slots

def cmd_slots(args):
    cfg = load_config(args.config)
    template = {s: [] for s in ["周六上午", "周六下午", "周日上午", "周日下午"]}
    if args.json:
        print(json.dumps(template, ensure_ascii=False, indent=2))
        return
    print(json.dumps(template, ensure_ascii=False, indent=2))


# ---------------------------------------------------------------- main

def main():
    p = argparse.ArgumentParser(description="质检部周工作自动化核心脚本")
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--json", action="store_true", help="以 JSON 输出，便于程序解析")
    common.add_argument("--config", help="指定配置文件路径")
    sub = p.add_subparsers(dest="cmd", required=True)

    p_info = sub.add_parser("info", parents=[common], help="查看表结构与活动清单")
    p_info.add_argument("xlsx")
    p_info.set_defaults(func=cmd_info)

    p_split = sub.add_parser("split", parents=[common], help="按校区拆分总表")
    p_split.add_argument("xlsx")
    p_split.add_argument("--outdir", help="输出目录")
    p_split.add_argument("--base-name", help="输出文件名主干")
    p_split.add_argument("--all", action="store_true",
                         help="按配置的校区清单出全部表；某校区本周无活动时出空白表并备注")
    p_split.set_defaults(func=cmd_split)

    p_arr = sub.add_parser("arrange", parents=[common], help="安排二级学院检查活动")
    p_arr.add_argument("xlsx")
    p_arr.add_argument("--week", type=int, help="周次（默认从标题提取）")
    p_arr.add_argument("--colleges", help="本周安排的学院，逗号分隔")
    p_arr.add_argument("--recommend", action="store_true", help="只输出候选推荐，不写文件")
    p_arr.add_argument("--dry-run", action="store_true", help="只预览不写文件")
    p_arr.add_argument("--force", action="store_true", help="忽略规则问题强制写入")
    p_arr.add_argument("--outdir")
    p_arr.add_argument("--base-name")
    p_arr.set_defaults(func=cmd_arrange)

    p_as = sub.add_parser("assistant", parents=[common], help="安排助理检查活动")
    p_as.add_argument("xlsx")
    p_as.add_argument("--week", type=int)
    p_as.add_argument("--slots", help="空闲时间表 json 文件")
    p_as.add_argument("--slots-json", help="直接传入 json 字符串")
    p_as.add_argument("--dry-run", action="store_true")
    p_as.add_argument("--outdir")
    p_as.add_argument("--base-name")
    p_as.set_defaults(func=cmd_assistant)

    p_chk = sub.add_parser("check", parents=[common], help="校验表格规则")
    p_chk.add_argument("xlsx")
    p_chk.add_argument("--week", type=int)
    p_chk.set_defaults(func=cmd_check)

    p_ing = sub.add_parser("ingest", parents=[common], help="把已排好的表读入台账（不修改表格）")
    p_ing.add_argument("xlsx")
    p_ing.add_argument("--week", type=int, help="周次（默认从标题/文件名提取）")
    p_ing.add_argument("--campus", help="校区（默认自动识别）")
    p_ing.set_defaults(func=cmd_ingest)

    p_lg = sub.add_parser("ledger", parents=[common], help="查看/清理台账")
    p_lg.add_argument("--week", type=int)
    p_lg.add_argument("--campus")
    p_lg.add_argument("--reset", action="store_true", help="清空台账（可配合 --week/--campus）")
    p_lg.set_defaults(func=cmd_ledger)

    p_sl = sub.add_parser("slots", parents=[common], help="生成空闲时间表模板")
    p_sl.set_defaults(func=cmd_slots)

    args = p.parse_args()
    try:
        args.func(args)
    except (ValueError, FileNotFoundError) as e:
        print(f"❌ {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
