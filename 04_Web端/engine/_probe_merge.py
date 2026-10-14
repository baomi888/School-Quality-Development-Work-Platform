# -*- coding: utf-8 -*-
"""
活动登记表汇总 · 回归门禁（纯 Python，秒级，非零退出即失败）

对应 README 里 `_probe_*.py` 的约定：任何一项不达标就返回非零退出码，可作为 CI 门禁。

覆盖：
  1. 跨周不混表（第4周 4 份 + 第5周 1 份 → 两套产物）
  2. 精确重复自动合并，溯源保留两个文件
  3. 表一（宣传部）剔除「活动对象=协会会员」
  4. 活动编号连续重排
  5. 排序规则：持续时间长→短；同日按周一→周日；同日按时间早晚
  6. 校验项：名称引号 / 级别别名 / 非法项目类别 / 全角标点 / 级别与对象不匹配
  7. 时间解析：跨周、跨天、同日多时段
  8. 无表头文件只报文件级错误，不抛异常
  9. 版式规格：标题宋体24加粗/行高30；正文宋体9/行高21；仅内容区有边框；备注红字
 10. 产出文件均可被 openpyxl 正常打开
"""
from __future__ import annotations

import shutil
import sys
import traceback
from pathlib import Path

import openpyxl
from openpyxl import Workbook

ENGINE = Path(__file__).resolve().parent
if str(ENGINE) not in sys.path:
    sys.path.insert(0, str(ENGINE))

from core import merge, zjb  # noqa: E402

FIX = ENGINE / "_probe_fixtures"
OUT = ENGINE / "_probe_out"

HDR_15 = ["活动编号", "活动名称", "活动级别", "主办单位", "开展时间", "地点", "负责人",
          "长号&短号", "活动类型", "项目类别", "活动对象", "参加方式", "备注", "校区", "加分分值"]

FAILS: list[str] = []
PASSES = [0]


def check(cond, label: str, detail: str = ""):
    if cond:
        PASSES[0] += 1
    else:
        FAILS.append(f"{label}{('  → ' + detail) if detail else ''}")


def eq(actual, expected, label: str):
    check(actual == expected, label, f"实际={actual!r} 期望={expected!r}")


def make_table(path: Path, week: int, rows, *, header=None, title=None,
               way_header="参加方式", junk_row=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    if junk_row:
        ws.cell(row=1, column=1, value=junk_row)
        title_row, hdr_row, first = 2, 3, 4
    else:
        title_row, hdr_row, first = 1, 2, 3
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=15)
    ws.cell(row=title_row, column=1, value=title or f"第{week}周活动登记表")
    hdr = list(header) if header else list(HDR_15)
    if not header:
        hdr[11] = way_header
    for i, h in enumerate(hdr, 1):
        ws.cell(row=hdr_row, column=i, value=h)
    for r, row in enumerate(rows, first):
        for i, v in enumerate(row, 1):
            ws.cell(row=r, column=i, value=v)
    wb.save(path)
    return path


# ------------------------------------------------------------------ 素材

A = ["", "2026年建筑工程学院大学生素质拓展中心结构设计大赛", "院级",
     "建筑工程学院大学生素质拓展中心", "周一09:00-周三00:00", "二教A201", "张三",
     "13800138000", "临时", "学术科技与专业技能", "全院学生", "线下", "", "西城", "0.2-1.0"]
B = ["", "2026年校青年志愿者协会会员培训", "校级", "校青年志愿者协会",
     "周日14:30-17:30", "图书馆报告厅", "李四", "13900139000", "常规",
     "社会实践与志愿服务", "协会会员", "线下", "", "官渡", "0.2"]
C = ["", "2026年体育学院篮球联赛", "院级", "体育学院学生会",
     "第三周周三-第五周周三", "篮球场", "王五", "13700137000", "临时",
     "文化艺术与体育竞技", "全院学生", "线下", "", "西城", "0.5"]
D = ["", '2026年外国语学院"英语角"活动', "院（系）级", "外国语学院学生会",
     "周三19:30-20:30", "线上", "赵六", "13600136000", "临时", "学术科技",
     "全校学生", "线上", "备注：含全角标点", "光华", "0.2"]
F = ["", "2026年建筑工程学院大学生素质拓展中心结构设计大赛", "院级",
     "建筑工程学院大学生素质拓展中心", "周五10:00-12:00", "二教A201", "张三",
     "13800138000", "临时", "学术科技与专业技能", "全院学生", "线下", "", "西城", "0.2-1.0"]
G = ["", "2026年理学院数学建模竞赛", "院级", "理学院学生会",
     "周六08:00-12:00", "实训楼302", "孙七", "13500135000", "临时",
     "学术科技与专业技能", "全院学生", "线下", "", "官渡", "0.2-0.8"]

_KEYS = [k for k, _ in merge.STD_COLUMNS]


def raw(row):
    """把 15 列的行按标准键装箱，供 build_record 直接使用"""
    return {k: v for k, v in zip(_KEYS, row)}


A_RAW, D_RAW = raw(A), raw(D)


def build_fixtures():
    if FIX.exists():
        shutil.rmtree(FIX)
    # 第4周
    make_table(FIX / "建工学院.xlsx", 4, [A])
    make_table(FIX / "化工学院.xlsx", 4, [A, B], way_header="参与方式", junk_row="广东石油化工学院")
    make_table(FIX / "外国语学院.xlsx", 4, [C, D])
    make_table(FIX / "体育学院.xlsx", 4, [F])
    # 第5周
    make_table(FIX / "理学院.xlsx", 5, [G])
    # 无表头
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="这不是活动登记表")
    ws.cell(row=2, column=1, value="随便什么内容")
    wb.save(FIX / "坏文件.xlsx")
    return FIX


# ------------------------------------------------------------------ 断言

def run():
    cfg = zjb.load_config(ENGINE / "core" / "config.json")
    build_fixtures()
    if OUT.exists():
        shutil.rmtree(OUT)

    result = merge.run_merge([FIX], OUT, cfg)

    # 1. 跨周不混表
    eq(result["week"], 4, "1. 主周次为第4周")
    eq(len(result["extra_weeks"]), 1, "1. 另有一套第5周产物")
    eq(result["extra_weeks"][0]["week"], 5, "1. 额外产物周次为5")

    # 8. 无表头文件只报文件级错误
    bad = [s for s in result["sources"] if s["file"] == "坏文件.xlsx"]
    eq(len(bad), 1, "8. 坏文件被读入并记录")
    eq(bad[0]["ok"], False, "8. 坏文件标记为失败")
    eq([i["code"] for i in bad[0]["issues"]], ["HEADER_NOT_FOUND"], "8. 报 HEADER_NOT_FOUND")

    # 2/4/5. 去重、编号、排序
    names = [a["name"] for a in result["activities"]]
    eq(len(names), 5, f"2. 第4周去重后 5 个活动（含 A 合并、F 疑似重复）")
    eq(result["counts"]["exact_merged"], 1, "2. 精确重复自动合并 1 条")
    eq(result["counts"]["unique"], 5, "2. unique 计数一致")
    expect_order = [C[1], A[1], D[1], F[1], B[1]]
    eq(names, expect_order, "5. 排序：15天 → 3天 → 1天(周三/周五/周日)")

    nos = [a["no"] for a in result["activities"]]
    eq(nos, ["1", "2", "3", "4", "5"], "4. 活动编号连续重排")

    a_rec = result["activities"][1]
    eq(len(a_rec["sources"]), 2, "2. A 的溯源保留了 2 个来源文件")
    eq(sorted(a_rec["sources"]), ["化工学院.xlsx", "建工学院.xlsx"], "2. A 的来源文件正确")

    # 7. 时间解析
    eq(result["activities"][0]["duration_days"], 15, "7. 跨周活动持续 15 天")
    eq(result["activities"][0]["slots"], ["第3周周三-第5周周三"], "7. 跨周时段标签完整")
    eq(result["activities"][1]["slots"], ["周一09:00-周三00:00"], "7. 跨天时段标签完整")
    eq(a_rec["duration_days"], 3, "7. 跨天活动持续 3 天")
    b_rec = result["activities"][4]
    eq(b_rec["slots"], ["周日14:30-17:30"], "7. 单时段标签")
    eq(b_rec["duration_days"], 1, "7. 单日活动持续 1 天")

    # 6. 校验项
    codes = {i["code"] for a in result["activities"] for i in a["issues"]}
    for code in ("NAME_QUOTE", "CATEGORY_INVALID", "FULLWIDTH_PUNCT",
                 "LEVEL_TARGET_MISMATCH", "PLACE_ONLINE_NO_LINK"):
        check(code in codes, f"6. 命中校验项 {code}")

    d_rec = [a for a in result["activities"] if a["name"].startswith("2026年外国语学院")][0]
    eq(d_rec["level"], "院级", "6. 级别别名「院（系）级」归一为「院级」")
    cat_issue = [i for i in d_rec["issues"] if i["code"] == "CATEGORY_INVALID"]
    check(cat_issue and cat_issue[0]["level"] == "error", "6. 非法项目类别是 error 级")

    # 6b. 协会会员 + 校级 是正确组合（培训 PPT：协会培训——协会会员），不得误报
    assoc_rec = [a for a in result["activities"] if "协会会员培训" in a["name"]][0]
    eq(assoc_rec["level"], "校级", "6b. 协会培训活动级别为校级")
    b_mismatch = [i for i in assoc_rec["issues"] if i["code"] == "LEVEL_TARGET_MISMATCH"]
    eq(len(b_mismatch), 0, "6b. 校级 + 协会会员不误报级别/对象不匹配")

    # 6c. 院级 + 协会会员 才应报
    assoc_bad = merge.build_record(
        {**A_RAW, "level": "院级", "target": "协会会员"}, cfg)
    check("LEVEL_TARGET_MISMATCH" in {i.code for i in assoc_bad.issues},
          "6c. 院级 + 协会会员 应报级别/对象不匹配")

    # 6b. 疑似重复
    eq(len(result["suspects"]), 1, "6. 同名不同时间 → 1 组疑似重复")
    eq(result["suspects"][0]["name"], A[1], "6. 疑似重复活动名正确")

    # 3. 表一剔除协会会员
    c = result["counts"]
    eq(c["unique"], 5, "3. 总表 5 行")
    eq(c["publicity_rows"], 4, "3. 表一 4 行（剔除协会会员）")
    pub_wb = openpyxl.load_workbook(result["outputs"]["publicity"])
    pub_targets = [pub_wb.active.cell(row=r, column=11).value
                   for r in range(3, pub_wb.active.max_row + 1)]
    check(all("协会会员" not in (t or "") for t in pub_targets), "3. 表一无协会会员活动")
    pub_wb.close()

    # 9. 版式规格
    wb = openpyxl.load_workbook(result["outputs"]["total"])
    ws = wb.active
    eq(ws["A1"].value, "第4周活动登记表", "9. 标题为「第4周活动登记表」")
    eq(ws["A1"].font.size, 24, "9. 标题字号 24")
    eq(bool(ws["A1"].font.bold), True, "9. 标题加粗")
    eq(ws["A1"].font.name, "宋体", "9. 标题字体宋体")
    eq(ws.row_dimensions[1].height, 30, "9. 标题行高 30 磅")
    eq(ws["B3"].font.size, 9, "9. 正文字号 9")
    eq(ws["B3"].font.name, "宋体", "9. 正文字体宋体")
    eq(ws.row_dimensions[3].height, 21, "9. 正文行高 21 磅")
    eq(ws["M3"].font.color.rgb, "FFFF0000", "9. 备注列红字")
    eq(ws.max_column, 15, "9. 总表 15 列")
    # 标题行外框（对齐官方《总表》做法：上/下通描，左在首格、右在末格）
    eq(ws["A1"].border.left.style, "thin", "9. 标题行左边框")
    eq(ws["A1"].border.top.style, "thin", "9. 标题行上边框")
    eq(ws["O1"].border.right.style, "thin", "9. 标题行右边框")
    eq(ws["B1"].border.left.style, None, "9. 标题行内部不画竖线")
    eq(ws["B1"].border.top.style, "thin", "9. 标题行中间格描上下边框")
    eq(ws["A1"].fill.fill_type, None, "9. 标题行无底色")
    eq(ws["A2"].fill.fill_type, None, "9. 表头无底色（对齐官方总表）")
    eq(ws["A30"].border.left.style, None, "9. 数据区之外没有边框")
    eq(ws["B30"].border.bottom.style, None, "9. 数据区之外没有下边框")
    borders_in_data = all(ws.cell(row=r, column=cc).border.left.style == "thin"
                          for r in range(2, 3 + c["unique"])
                          for cc in range(1, 16))
    check(borders_in_data, "9. 内容区每格都有边框")
    # 列宽按官方总表标定（宋体9号中文双宽，窄列会让长名折行后被21磅行高裁掉）
    eq(round(ws.column_dimensions["B"].width, 2), 52.21, "9. 活动名称列宽 52.21（防长名折行）")
    names_len = max(len(a["name"]) for a in result["activities"])
    check(names_len >= 24, "9. 素材含长活动名用于折行校验", f"最长 {names_len} 字")
    wb.close()

    # 9b. 超长内容 → CELL_WRAP_CLIP 提醒（会折行、又被固定行高裁掉）
    over = dict(D_RAW)
    over["name"] = "2026年" + "建" * 40 + "学院大学生素质拓展中心超长活动名称占位测试"
    over_rec = merge.build_record(over, cfg)
    over_codes = {i.code for i in over_rec.issues}
    check("CELL_WRAP_CLIP" in over_codes, "9b. 超长活动名称命中 CELL_WRAP_CLIP")
    # 正常长度的真实素材不应误报
    ok_rec = merge.build_record(dict(A_RAW), cfg)
    check("CELL_WRAP_CLIP" not in {i.code for i in ok_rec.issues},
          "9b. 正常长度活动名不误报 CELL_WRAP_CLIP")

    # 10. 四个产物都能打开
    for key, path in result["outputs"].items():
        check(Path(path).exists(), f"10. 产物存在：{key}")
        try:
            w2 = openpyxl.load_workbook(path)
            w2.close()
            PASSES[0] += 1
        except Exception as exc:  # noqa: BLE001
            FAILS.append(f"10. 产物无法打开：{key} → {exc}")

    # 时间表行数 = 各活动时段数之和
    twb = openpyxl.load_workbook(result["outputs"]["timetable"])
    trows = twb.active.max_row - 2
    twb.close()
    eq(trows, 5, "7. 时间表 5 行（每个活动恰好 1 个时段）")

    # 指定 --week 时排除他周文件
    r5 = merge.run_merge([FIX], OUT / "w5only", cfg, 5)
    eq(r5["week"], 5, "1. 指定第5周时主周次为 5")
    eq(len(r5["dropped_files"]), 4, "1. 指定第5周时排除 4 个第4周文件")
    check(all(d["reason"] for d in r5["dropped_files"]), "1. 每个被排除文件都有原因")


def main() -> int:
    try:
        run()
    except Exception:  # noqa: BLE001
        traceback.print_exc()
        FAILS.append("probe 运行时异常（见上方 traceback）")

    print(f"\n活动登记表汇总门禁：通过 {PASSES[0]} 项，失败 {len(FAILS)} 项")
    if FAILS:
        print("\n失败项：")
        for f in FAILS:
            print("  ✗", f)
        return 1
    print("  ✓ 全部通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
