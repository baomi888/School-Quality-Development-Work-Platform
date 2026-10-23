# -*- coding: utf-8 -*-
"""
质检部周工作自动化 · Web 引擎

以 FastAPI 暴露能力，核心逻辑全部复用 core/zjb.py（与 CLI 技能同源，不重复实现）。
- 上传 / 解析登记表
- 按校区拆分
- 安排二级学院（候选推荐 + 规则校验 + 写入）
- 安排助理
- 台账入档 / 查询 / 清理 / 统计
- 产物下载（保留原表字体、边框、合并单元格）

后期迁移到服务器：把本目录整体部署即可，前端只需把 API 地址换成服务器地址。
"""
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.request
import uuid
import math
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import Optional, Dict, Any, List

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))
import core.zjb as zjb  # noqa: E402
import core.merge as mergemod  # noqa: E402
import core.planner as planner  # noqa: E402
import core.agent as agentmod  # noqa: E402

WORK_DIR = BASE / "work"
OUT_DIR = BASE / "output"
DATA_DIR = BASE / "data"
PREVIEW_DIR = BASE / "preview"
REGISTRY = DATA_DIR / "registry.json"
WEB_DIST = BASE / "web_dist"
# 「活动登记表汇总」工具页：零构建的独立页面，不依赖 web_dist（前端重建失败也能用）
MERGE_TOOL = BASE / "merge_tool.html"
# OfficeCLI 为可选渲染器（用于高保真预览），缺失时自动回退内置渲染
OFFICECLI = BASE / "bin" / ("officecli.exe" if os.name == "nt" else "officecli")

for d in (WORK_DIR, OUT_DIR, DATA_DIR, PREVIEW_DIR):
    d.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------- 配置与台账

def get_cfg() -> dict:
    """加载配置；台账默认与 CLI 技能共用同一份（单一事实源），可用环境变量覆盖"""
    cfg = zjb.load_config(str(BASE / "core" / "config.json"))
    env_ledger = os.environ.get("ZJB_LEDGER")
    if env_ledger:
        cfg["ledger_path"] = env_ledger
    elif (BASE / ".portable").exists():
        # 便携模式（免安装包）：只认本目录的台账，绝不读开发机上的共享技能目录
        cfg["ledger_path"] = str(DATA_DIR / "ledger.json")
    else:
        shared = Path.home() / ".workbuddy" / "skills" / "zhijianbu-weekly" / "data" / "ledger.json"
        cfg["ledger_path"] = str(shared if shared.parent.exists() else (DATA_DIR / "ledger.json"))
    return cfg


def cfg_for_campus(cfg: dict, campus: Optional[str]) -> dict:
    cfg = dict(cfg)
    cfg["_current_campus"] = campus
    return cfg


# ---------------------------------------------------------------- 文件登记表

def load_registry() -> dict:
    if REGISTRY.exists():
        try:
            return json.loads(REGISTRY.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {}
    return {}


def save_registry(reg: dict):
    REGISTRY.write_text(json.dumps(reg, ensure_ascii=False, indent=2), encoding="utf-8")


def register(path: Path, kind: str, extra: Optional[dict] = None) -> dict:
    reg = load_registry()
    fid = uuid.uuid4().hex[:12]
    reg[fid] = {
        "path": str(path),
        "name": path.name,
        "kind": kind,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        **(extra or {}),
    }
    save_registry(reg)
    return {"id": fid, "name": path.name, "kind": kind, **(extra or {})}


def resolve(fid: str) -> Path:
    reg = load_registry()
    item = reg.get(fid)
    if not item or not Path(item["path"]).exists():
        raise HTTPException(404, "文件不存在或已清理，请重新上传")
    return Path(item["path"])


# ---------------------------------------------------------------- 业务封装

def open_sheet(path: Path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header_row, header_map = zjb.find_header_row(ws)
    if header_row is None:
        raise HTTPException(400, "没找到表头行，请确认是标准的活动登记表（第 2 行为表头）")
    return wb, ws, header_row, header_map


def guess_week(cfg, ws, src: Path, week: Optional[int]):
    if week:
        return int(week)
    try:
        title = ws.cell(row=1, column=1).value
    except Exception:
        title = None
    return zjb.extract_week(title) or zjb.extract_week(src.name)


def summarize(wb, ws, header_row, header_map, cfg, src: Path, week) -> dict:
    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    campuses = sorted({a["campus"] for a in acts if a["campus"]})
    _, checker_map = zjb.read_checkers(ws, header_row, cfg["checker_column"], cfg)
    _, assistant_map = zjb.read_checkers(ws, header_row, cfg["assistant_column"], cfg)
    return {
        "week": week,
        "campuses": campuses,
        "activity_count": len(acts),
        "checker_records": sum(len(v) for v in checker_map.values()),
        "assistant_records": sum(len(v) for v in assistant_map.values()),
        "has_checkers": bool(checker_map),
        "has_assistants": bool(assistant_map),
        "activities": [
            {
                "row": a["row"], "no": a.get("no"), "name": a.get("name"),
                "host_college": a.get("host_college"), "time": a.get("time"),
                "campus": a.get("campus"), "place": a.get("place"),
                "level": a.get("level"),
            } for a in acts
        ],
        "header_map": {k: v for k, v in header_map.items()},
    }


def do_split(cfg, src: Path, outdir: Path, all_campuses: bool) -> dict:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = zjb.find_header_row(ws)
    campus_col = header_map.get("campus")
    if not campus_col:
        raise HTTPException(400, "表中没有「校区」列，无法拆分")
    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    if not acts:
        raise HTTPException(400, "表里没有活动数据行，无法拆分")
    week = guess_week(cfg, ws, src, None)

    groups: Dict[str, List[int]] = {}
    for a in acts:
        groups.setdefault(a["campus"] or "未知校区", []).append(a["row"])

    known = zjb.campus_in_name(src, cfg)
    if known and not all_campuses:
        targets = [known]
    else:
        targets = list(dict.fromkeys((cfg.get("campuses") or []) + list(groups.keys())))

    outdir.mkdir(parents=True, exist_ok=True)
    head_rows = list(range(1, header_row + 1))
    last_col = max(header_map.values())
    base = zjb.standard_base_name(cfg, week)

    outputs = []
    for campus in targets:
        rows = sorted(groups.get(campus, []))
        wb2, ws2 = zjb.clone_subset(ws, head_rows + rows, last_col=last_col)
        empty = not rows
        if empty:
            zjb.fill_empty_week(ws2, ws, header_row, last_col, header_map, cfg)
        out = zjb.out_path_for(src, f"（{campus}）", base_name=base, outdir=str(outdir), cfg=cfg)
        wb2.save(out)
        f = register(out, "split", {"campus": campus, "week": week, "empty": empty, "rows": len(rows)})
        outputs.append({**f, "campus": campus, "rows": len(rows), "empty": empty})
    return {"week": week, "outputs": outputs}


def do_recommend(cfg, src: Path, week) -> dict:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = zjb.find_header_row(ws)
    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    if not acts:
        raise HTTPException(400, "表中没有读到任何活动行")
    week = guess_week(cfg, ws, src, week)
    campus = zjb.normalize_campus(acts[0].get("campus"), cfg) or "未知"
    cfg2 = cfg_for_campus(cfg, campus)
    ledger = zjb.load_ledger(cfg2)
    return zjb.recommend_pool(acts, cfg2, ledger, week)


def do_arrange(cfg, src: Path, week, colleges: List[str], outdir: Path, force: bool, dry_run: bool = False) -> dict:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = zjb.find_header_row(ws)
    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    if not acts:
        raise HTTPException(400, "表中没有读到任何活动行")
    week = guess_week(cfg, ws, src, week)
    campus = zjb.normalize_campus(acts[0].get("campus"), cfg) or "未知"
    cfg2 = cfg_for_campus(cfg, campus)
    ledger = zjb.load_ledger(cfg2)

    pool = [c.strip() for c in colleges if c and c.strip()]
    if not pool:
        raise HTTPException(400, "请至少选择一个学院")

    last, _ = zjb.last_week_checkers(ledger, week, cfg2)
    per = cfg["per_activity_colleges"]
    # 排序链抽在 core/planner.py：智能体工具与 Web 接口共用同一份确定性实现
    plan, load, problems = planner.plan_checkers(acts, pool, last, per)

    if problems and not force:
        alt = [c for c in cfg["colleges"] if c not in last and c not in pool]
        raise HTTPException(409, {
            "message": "规则校验未通过，已停止写入",
            "problems": problems,
            "alternatives": alt,
        })

    col_idx, is_new = zjb.ensure_checker_columns(ws, header_row, header_map, per, cfg)
    for row_idx, names in plan.items():
        for i, name in enumerate(names):
            cell = ws.cell(row=row_idx, column=col_idx + i)
            cell.value = name
            cell._style = copy(ws.cell(row=row_idx, column=1)._style)
    if is_new:
        zjb.extend_title_merge(ws, col_idx + per - 1)

    outdir.mkdir(parents=True, exist_ok=True)
    out = zjb.out_path_for(src, cfg["suffix_colleges"], outdir=str(outdir), cfg=cfg)
    written = not dry_run
    if written:
        wb.save(out)
        zjb.record_to_ledger(ledger, cfg2, week, campus, acts, plan, "colleges")
        zjb.save_ledger(cfg2, ledger)

    result = {
        "week": week, "campus": campus, "written": written,
        "column": openpyxl.utils.get_column_letter(col_idx),
        "plan": [{"row": a["row"], "no": a.get("no"), "name": a.get("name"),
                  "host_college": a.get("host_college"), "checkers": plan.get(a["row"], [])} for a in acts],
        "pool": pool, "load": load,
        "last_week_checkers": sorted(last),
        "problems": problems,
    }
    if written:
        result["file"] = register(out, "arranged", {"campus": campus, "week": week})
    return result


def do_assistant(cfg, src: Path, week, slots: dict, outdir: Path) -> dict:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = zjb.find_header_row(ws)
    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    week = guess_week(cfg, ws, src, week)
    campus = zjb.normalize_campus(acts[0].get("campus"), cfg) if acts else "未知"

    per = cfg["per_activity_assistants"]
    load: Dict[str, int] = {}
    plan: Dict[int, List[str]] = {}
    problems: List[str] = []
    for a in acts:
        need = zjb.parse_slots(a["time"])
        if not need:
            problems.append(f"活动「{a['name']}」时间「{a['time']}」不是周末时段，跳过")
            continue
        sets = [set(slots.get(s, [])) for s in need]
        avail = set.intersection(*sets) if sets else set()
        loose = False
        if len(avail) < per:
            union = set.union(*sets) if sets else set()
            if len(union) >= per:
                avail, loose = union, True
        avail_sorted = sorted(avail, key=lambda p: (load.get(p, 0), p))
        chosen = avail_sorted[:per]
        for p in chosen:
            load[p] = load.get(p, 0) + 1
        plan[a["row"]] = chosen
        if len(avail) < per:
            problems.append(f"活动「{a['name']}」可用助理不足（需{per}人，时段={'+'.join(need)}）")
        elif loose and chosen:
            problems.append(f"活动「{a['name']}」按宽松匹配安排（无人全程有空）")

    col_idx, is_new = zjb.ensure_column(ws, header_row, header_map,
                                        cfg["assistant_column"], cfg["assistant_column"][0], cfg)
    zjb.write_values(ws, col_idx, plan, cfg)
    if is_new:
        zjb.extend_title_merge(ws, col_idx)
    outdir.mkdir(parents=True, exist_ok=True)
    out = zjb.out_path_for(src, cfg["suffix_assistants"], outdir=str(outdir), cfg=cfg)
    wb.save(out)

    return {
        "week": week, "campus": campus, "written": True,
        "column": openpyxl.utils.get_column_letter(col_idx),
        "plan": [{"row": a["row"], "no": a.get("no"), "name": a.get("name"),
                  "slots": zjb.parse_slots(a["time"]),
                  "assistants": plan.get(a["row"], [])} for a in acts],
        "load": load, "problems": problems,
        "file": register(out, "assistant", {"campus": campus, "week": week}),
    }


def do_check(cfg, src: Path, week) -> dict:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = zjb.find_header_row(ws)
    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    week = guess_week(cfg, ws, src, week)
    campus = zjb.normalize_campus(acts[0].get("campus"), cfg) if acts else "未知"
    cfg2 = cfg_for_campus(cfg, campus)
    col, checker_map = zjb.read_checkers(ws, header_row, cfg["checker_column"], cfg)
    ledger = zjb.load_ledger(cfg2)
    last, _ = zjb.last_week_checkers(ledger, week, cfg2)
    issues = []
    for a in acts:
        names = checker_map.get(a["row"], [])
        if len(names) != cfg["per_activity_colleges"]:
            issues.append({"row": a["row"], "name": a.get("name"), "type": "数量不符",
                           "detail": f"应为{cfg['per_activity_colleges']}个学院，实际{len(names)}个"})
        if a["host_college"] and a["host_college"] in names:
            issues.append({"row": a["row"], "name": a.get("name"), "type": "自检",
                           "detail": f"主办学院 {a['host_college']} 不能检查自己"})
        dup = [n for n in names if n in last]
        if dup:
            issues.append({"row": a["row"], "name": a.get("name"), "type": "连续两周",
                           "detail": "上周已安排：" + "、".join(dup)})
    return {"week": week, "campus": campus, "ok": not issues, "issues": issues,
            "checker_column": openpyxl.utils.get_column_letter(col) if col else None}


def do_ingest(cfg, src: Path, week, campus: Optional[str]) -> dict:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = zjb.find_header_row(ws)
    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    week = guess_week(cfg, ws, src, week)
    if not campus:
        campus = zjb.normalize_campus(acts[0].get("campus"), cfg) if acts else None
    if not campus:
        m = re.search(r"[（(](官渡|西城)[）)]", src.name)
        campus = m.group(1) if m else None
    if not campus:
        raise HTTPException(400, "无法确定校区，请手动指定")
    cfg2 = cfg_for_campus(cfg, campus)
    c_col, checker_map = zjb.read_checkers(ws, header_row, cfg["checker_column"], cfg)
    a_col, asst_map = zjb.read_checkers(ws, header_row, cfg["assistant_column"], cfg)
    if not checker_map and not asst_map:
        raise HTTPException(400, "表中没有找到检查学院/检查助理数据")
    ledger = zjb.load_ledger(cfg2)
    if checker_map:
        zjb.record_to_ledger(ledger, cfg2, week, campus, acts, checker_map, "colleges")
    if asst_map:
        zjb.record_to_ledger(ledger, cfg2, week, campus, acts, asst_map, "assistants")
    zjb.save_ledger(cfg2, ledger)
    return {"week": week, "campus": campus, "ingested": True,
            "activities": [{"no": a.get("no"), "name": a.get("name"),
                            "checkers": checker_map.get(a["row"]) or [],
                            "assistants": asst_map.get(a["row"]) or []} for a in acts]}


def do_import_history(cfg, src: Path, week, force: bool = False) -> dict:
    """导入历史活动登记表：不重排。
    - 有安排列：读取既有检查学院/助理，补录完整台账
    - 无安排列：作为共享原表归档，并只记录活动；不会伪造检查或助理
    - 支持单校区表与多校区总表（自动按校区分组后分别入账）
    - 周次从文件名/表格标题识别，也可通过参数指定
    - 原文件归档到 data/history/第X周/ 长期留存
    """
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    header_row, header_map = zjb.find_header_row(ws)
    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    if not acts:
        raise HTTPException(400, "表中没有活动数据行，无法导入")
    week = guess_week(cfg, ws, src, week)
    if not week:
        raise HTTPException(400, "无法识别周次，请在参数里指定 week")

    c_col, checker_map = zjb.read_checkers(ws, header_row, cfg["checker_column"], cfg)
    a_col, asst_map = zjb.read_checkers(ws, header_row, cfg["assistant_column"], cfg)
    has_arrangements = bool(checker_map or asst_map)

    # 按校区分组；校区列缺失时尝试从文件名推断
    groups: Dict[str, List[dict]] = {}
    for a in acts:
        campus = zjb.normalize_campus(a.get("campus"), cfg) or "未知"
        groups.setdefault(campus, []).append(a)
    if list(groups.keys()) == ["未知"]:
        known = zjb.campus_in_name(src, cfg)
        if known:
            groups = {known: acts}

    # 归档原档（长期留存）—— 尽力而为：目标被占用或已存在时换名/跳过，不影响台账补录
    hist_dir = DATA_DIR / "history" / f"第{week}周"
    hist_dir.mkdir(parents=True, exist_ok=True)
    dest_hist = hist_dir / src.name
    try:
        shutil.copy2(src, dest_hist)
    except OSError:
        try:
            ts = datetime.now().strftime("%H%M%S")
            dest_hist = hist_dir / f"{dest_hist.stem}_{ts}{dest_hist.suffix}"
            shutil.copy2(src, dest_hist)
        except OSError:
            dest_hist = None

    ledger = zjb.load_ledger(cfg)
    campuses_out = []
    warnings: List[str] = []
    for campus, cacts in groups.items():
        cfg2 = cfg_for_campus(cfg, campus)
        rows = {a["row"] for a in cacts}
        ck = sum(len(checker_map.get(r, [])) for r in rows)
        ast = sum(len(asst_map.get(r, [])) for r in rows)
        # 即使没有安排列，也保存活动本身；空安排不会计入学院轮值或助理统计。
        key = zjb.record_to_ledger(
            ledger, cfg2, week, campus, cacts,
            {a["row"]: [] for a in cacts}, "activities")
        if checker_map:
            zjb.record_to_ledger(ledger, cfg2, week, campus, cacts, checker_map, "colleges")
        if asst_map:
            zjb.record_to_ledger(ledger, cfg2, week, campus, cacts, asst_map, "assistants")
        entry = ledger["weeks"][key]
        has_existing = bool(
            entry.get("checkers")
            or any(a.get("assistants") for a in entry.get("activities", []))
        )
        entry["source_status"] = (
            "assigned" if checker_map and asst_map
            else "partially-assigned" if has_arrangements
            else "existing-assignments" if has_existing
            else "activities-only"
        )
        campuses_out.append({"campus": campus, "activities": len(cacts),
                             "checker_records": ck, "assistant_records": ast,
                             "checker_rows": len([a for a in cacts if checker_map.get(a["row"])]),
                             "assistant_rows": len([a for a in cacts if asst_map.get(a["row"])])})
    if not checker_map:
        warnings.append("未找到检查学院列，本次只记录活动，不计入学院轮值。")
    elif any(c["checker_rows"] < c["activities"] for c in campuses_out):
        warnings.append("部分活动没有检查学院，未填写的活动会保留为空。")
    if not asst_map:
        warnings.append("未找到检查助理列，助理安排会保留为空。")
    elif any(c["assistant_rows"] < c["activities"] for c in campuses_out):
        warnings.append("部分活动没有检查助理，未填写的活动会保留为空。")
    zjb.save_ledger(cfg, ledger)

    return {
        "week": week,
        "mode": "ledger" if has_arrangements else "activities_only",
        "arrangements_imported": has_arrangements,
        "ledger_updated": True,
        "warnings": list(dict.fromkeys(warnings)),
        "archived_to": str(dest_hist) if dest_hist else "（归档跳过：目标文件被占用）",
        "campuses": campuses_out,
        "total_activities": sum(len(v) for v in groups.values()),
    }


# ---------------------------------------------------------------- 统计

def do_sync_assistants(cfg, source_week: int, target_week: int, task: str,
                        plan_url: str = "http://127.0.0.1:8766/api/plan") -> dict:
    """联动：从 05 工作安排台取第 source_week 周「安排助理检查活动」的西城/官渡助理名单，
    按校区写回 04 第 target_week 周台账各活动的 assistants 字段。

    场景：第 7 周（在 05 安排）处理的是第 6 周开展的活动（在 04 台账），
    于是第 6 周每个活动上要显示第 7 周被安排去检查的助理。"""
    url = f"{plan_url}?dept=zjb&week={source_week}"
    try:
        with urllib.request.urlopen(url, timeout=6) as resp:
            data = json.loads(resp.read().decode("utf-8"))
    except Exception as e:
        raise HTTPException(502, f"连不上工作安排台（05，端口 8766）：{e}。请先启动工作安排台再同步。")
    grid = data.get("assign") or {}
    task_cells = grid.get(task) or {}
    xc = task_cells.get("西城", []) or []
    gd = task_cells.get("官渡", []) or []
    if not xc and not gd:
        return {"ok": True, "synced": 0,
                "note": f"第{source_week}周计划中「{task}」没有西城/官渡助理，未写入"}
    ledger = zjb.load_ledger(cfg)
    weeks = ledger.get("weeks", {})
    target_keys = [k for k in weeks if k.startswith(f"W{target_week}-")]
    if not target_keys:
        raise HTTPException(404, f"04 台账里没有第{target_week}周记录（请先导入第{target_week}周活动登记表）")
    synced = 0
    for key in target_keys:
        campus = weeks[key].get("campus")
        names = gd if campus == "官渡" else xc
        if not names:
            continue
        for a in weeks[key].get("activities", []):
            a["assistants"] = list(names)
            synced += 1
    zjb.save_ledger(cfg, ledger)
    return {"ok": True, "synced": synced, "target_week": target_week,
            "西城助理": xc, "官渡助理": gd, "weeks": target_keys}



def build_stats(cfg) -> dict:
    ledger = zjb.load_ledger(cfg)
    weeks = ledger.get("weeks", {})
    college_count: Dict[str, int] = {}
    per_week: Dict[str, dict] = {}
    campus_count: Dict[str, int] = {}

    for key, rec in weeks.items():
        w = rec.get("week")
        campus = rec.get("campus") or "-"
        label = f"第{w}周" if w else key
        bucket = per_week.setdefault(label, {"week": w or 0, "label": label,
                                             "activities": 0, "colleges": set(), "campuses": set()})
        for a in rec.get("activities", []):
            bucket["activities"] += 1
            campus_count[a.get("campus") or campus] = campus_count.get(a.get("campus") or campus, 0) + 1
            for c in a.get("checkers") or []:
                college_count[c] = college_count.get(c, 0) + 1
                bucket["colleges"].add(c)
            bucket["campuses"].add(a.get("campus") or campus)

    per_week_list = [{
        "label": v["label"], "week": v["week"],
        "activities": v["activities"],
        "colleges": len(v["colleges"]),
    } for v in sorted(per_week.values(), key=lambda x: x["week"])]

    return {
        "college_load": sorted(
            [{"college": k, "count": v} for k, v in college_count.items()],
            key=lambda x: -x["count"]),
        "per_week": per_week_list,
        "campus_distribution": [{"campus": k, "count": v} for k, v in campus_count.items()],
        "total_weeks": len(weeks),
        "total_activities": sum(v["activities"] for v in per_week.values()),
        "college_pool": cfg.get("colleges", []),
        "ledger_file": str(zjb.ledger_file(cfg)),
    }


# ---------------------------------------------------------------- API 模型

class ArrangeIn(BaseModel):
    file_id: str
    week: Optional[int] = None
    colleges: List[str] = []
    force: bool = False
    dry_run: bool = False


class AssistantIn(BaseModel):
    file_id: str
    week: Optional[int] = None
    slots: Dict[str, List[str]] = {}


class SplitIn(BaseModel):
    file_id: str
    all_campuses: bool = True


class SimpleIn(BaseModel):
    file_id: str
    week: Optional[int] = None
    campus: Optional[str] = None


class ImportHistoryIn(BaseModel):
    file_id: str
    week: Optional[int] = None
    force: bool = False


class SyncAssistIn(BaseModel):
    source_week: int = 7
    target_week: int = 6
    task: str = "t_03"
    plan_url: str = "http://127.0.0.1:8766/api/plan"


# ---------------------------------------------------------------- App

app = FastAPI(title="质检部周工作自动化", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/api/health")
def health():
    cfg = get_cfg()
    return {"ok": True, "term": cfg.get("term_name"),
            "colleges": cfg.get("colleges", []),
            "campuses": cfg.get("campuses", []),
            "ledger_file": str(zjb.ledger_file(cfg))}


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "只支持 .xlsx / .xlsm；旧版 .xls 请先用 Excel 或 WPS 另存为 .xlsx")
    cfg = get_cfg()
    folder = WORK_DIR / uuid.uuid4().hex[:12]
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|]', "_", file.filename)
    dest = folder / safe
    with dest.open("wb") as f:
        f.write(await file.read())

    try:
        wb, ws, header_row, header_map = open_sheet(dest)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            400,
            f"无法读取这个 Excel：{exc}。请确认它是可正常打开的 .xlsx，并可尝试在 Excel / WPS 中另存为 .xlsx 后重试。",
        )
    week = guess_week(cfg, ws, dest, None)
    info = summarize(wb, ws, header_row, header_map, cfg, dest, week)
    reg = register(dest, "upload", {"week": week})
    return {"file": reg, "info": info}


@app.get("/api/files/{fid}/info")
def file_info(fid: str):
    cfg = get_cfg()
    src = resolve(fid)
    wb, ws, header_row, header_map = open_sheet(src)
    week = guess_week(cfg, ws, src, None)
    info = summarize(wb, ws, header_row, header_map, cfg, src, week)
    return {"file_id": fid, "name": src.name, "info": info}


@app.post("/api/split")
def split(body: SplitIn):
    cfg = get_cfg()
    src = resolve(body.file_id)
    outdir = OUT_DIR / uuid.uuid4().hex[:8]
    return do_split(cfg, src, outdir, body.all_campuses)


# ------------------------------------------------- 活动登记表汇总（规划部）
# 现有链路是「一张总表 → 按校区拆开 → 安排二级学院/助理」；
# 这里是它的反向：把各学院/各校区交上来的 N 份单表合成一张总表。

MERGE_LABELS = {
    "total": "总表（表二 · 质检部）",
    "publicity": "表一（宣传部）",
    "timetable": "时间表",
    "report": "汇总校验报告",
}


@app.post("/api/merge")
async def merge_tables(files: List[UploadFile] = File(...), week: Optional[int] = None):
    """多份单表 → 总表 / 表一（宣传部）/ 时间表 / 校验报告。

    规划部过去逐个学院下载再手工拼表，这里一次全部丢进来即可。
    不做文件名假设：来源按「内容里的校区列 + 文件名线索」双重判定；
    输入混了多个周次时按周分组、各出一套产物，绝不跨周混表。
    """
    if not files:
        raise HTTPException(400, "请至少上传一份活动登记表")
    cfg = get_cfg()
    folder = WORK_DIR / ("merge_" + uuid.uuid4().hex[:10])
    folder.mkdir(parents=True, exist_ok=True)
    for f in files:
        name = Path(f.filename or "未命名.xlsx").name
        if not name.lower().endswith((".xlsx", ".xlsm")):
            raise HTTPException(400, f"只支持 .xlsx / .xlsm：{name}（.xls 请先另存为 .xlsx）")
        safe = re.sub(r'[\\/:*?"<>|]', "_", name)
        (folder / safe).write_bytes(await f.read())

    outdir = OUT_DIR / uuid.uuid4().hex[:8]
    try:
        result = mergemod.run_merge([folder], outdir, cfg, week)
    except ValueError as exc:
        raise HTTPException(400, str(exc))

    outputs = {}
    for key, path in (result.get("outputs") or {}).items():
        p = Path(path)
        if not p.exists():
            continue
        reg = register(p, "merge-" + key, {"week": result.get("week")})
        outputs[key] = {
            "key": key,
            "label": MERGE_LABELS.get(key, key),
            "file_id": reg["id"],
            "name": p.name,
            "size": p.stat().st_size,
            "url": f"/api/files/{reg['id']}/download",
        }
    result["outputs"] = outputs

    # 额外周次的产物也要能下载
    for extra in result.get("extra_weeks") or []:
        outs = {}
        for key, path in (extra.get("outputs") or {}).items():
            p = Path(path)
            if not p.exists():
                continue
            reg = register(p, "merge-" + key, {"week": extra.get("week")})
            outs[key] = {
                "key": key, "label": MERGE_LABELS.get(key, key), "file_id": reg["id"],
                "name": p.name, "size": p.stat().st_size,
                "url": f"/api/files/{reg['id']}/download",
            }
        extra["outputs"] = outs
    return result


@app.get("/merge")
def merge_page():
    """零构建的汇总工具页（不依赖 web_dist 构建产物）"""
    if not MERGE_TOOL.exists():
        raise HTTPException(404, "汇总工具页 merge_tool.html 缺失")
    return HTMLResponse(MERGE_TOOL.read_text(encoding="utf-8"))


# ------------------------------------------------- 新手引导
# 免安装包里引导页放在包根（本文件的上一级目录的上一级）；
# 开发环境回落到项目根的 打包模板/，方便直接在 8765 预览。
GUIDE_CANDIDATES = [
    BASE.parent / "00-新手引导.html",                        # 包内：包根/00-新手引导.html
    BASE.parent / "新手引导.html",                           # 兼容旧包名
    BASE.parent.parent / "打包模板" / "00-新手引导.html",     # 开发环境预览
]


@app.get("/guide")
def guide_page():
    """新手引导页：启动后自动打开的第一个页面。"""
    for p in GUIDE_CANDIDATES:
        if p.exists():
            return HTMLResponse(p.read_text(encoding="utf-8"))
    raise HTTPException(404, "新手引导页缺失（00-新手引导.html）")


# ------------------------------------------------- 智能体（编排层）
# 网页只做"提问 + 展示执行轨迹"，规划与裁决都在 core/agent.py 里。
# 未配置 LLM 时自动降级为确定性规划器，所以这个演示页离线也能跑。

AGENT_DEMO = BASE / "agent_demo.html"


class AgentIn(BaseModel):
    goal: str
    inbox: Optional[str] = None
    week: Optional[int] = None
    table: Optional[str] = None
    campus: Optional[str] = None
    pool: Optional[str] = None
    doc: Optional[str] = None


@app.post("/api/agent")
def agent_run(body: AgentIn):
    """跑一次智能体闭环，返回完整执行轨迹（感知/规划/行动/反思）。"""
    ctx = {k: v for k, v in {
        "inbox": body.inbox, "week": body.week, "table": body.table,
        "campus": body.campus, "pool": body.pool, "doc": body.doc,
    }.items() if v not in (None, "")}
    try:
        return agentmod.run(body.goal, **ctx)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(500, f"智能体运行失败：{type(exc).__name__}: {exc}")


@app.get("/agent")
def agent_page():
    if not AGENT_DEMO.exists():
        raise HTTPException(404, "智能体演示页 agent_demo.html 缺失")
    return HTMLResponse(AGENT_DEMO.read_text(encoding="utf-8"))


@app.get("/api/agent/presets")
def agent_presets():
    """演示页的目标预设：把"能说一句话就让智能体干活"的场景摆出来。"""
    cfg = get_cfg()
    return {
        "presets": [
            {"goal": "把本周登记表汇总一下并生成简报",
             "hint": "收件目录 → 去重合并 → 校验 → 起草简报",
             "fields": ["inbox", "week"]},
            {"goal": "哪些学院还没交表",
             "hint": "扫描收件 → 对照应交单位算差集 → 生成催报通知",
             "fields": ["inbox"]},
            {"goal": "学分办法里比赛类怎么算分",
             "hint": "在规范性文件原文里检索，返回原文摘录", "fields": []},
            {"goal": "安排二级学院检查",
             "hint": "对校区表跑确定性排序链，只试算不写回",
             "fields": ["table", "campus"]},
            {"goal": "预审这份申报书",
             "hint": "按硬规则逐条预审，输出必改项 / 通过 / 无法判断",
             "fields": ["doc"]},
        ],
        "colleges": cfg.get("colleges", []),
        "campuses": cfg.get("campuses", []),
    }


@app.get("/api/recommend")
def recommend(file_id: str, week: Optional[int] = None):
    cfg = get_cfg()
    return do_recommend(cfg, resolve(file_id), week)


@app.post("/api/arrange")
def arrange(body: ArrangeIn):
    cfg = get_cfg()
    src = resolve(body.file_id)
    outdir = OUT_DIR / uuid.uuid4().hex[:8]
    return do_arrange(cfg, src, body.week, body.colleges, outdir, body.force, body.dry_run)


@app.post("/api/assistant")
def assistant(body: AssistantIn):
    cfg = get_cfg()
    src = resolve(body.file_id)
    outdir = OUT_DIR / uuid.uuid4().hex[:8]
    return do_assistant(cfg, src, body.week, body.slots, outdir)


@app.post("/api/check")
def check(body: SimpleIn):
    cfg = get_cfg()
    return do_check(cfg, resolve(body.file_id), body.week)


@app.post("/api/ingest")
def ingest(body: SimpleIn):
    cfg = get_cfg()
    return do_ingest(cfg, resolve(body.file_id), body.week, body.campus)


@app.post("/api/import-history")
def import_history(body: ImportHistoryIn):
    """导入历史活动登记表并补录台账（不重排、不另存表、归档原档）。"""
    cfg = get_cfg()
    src = resolve(body.file_id)
    try:
        return do_import_history(cfg, src, body.week, body.force)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"历史表导入失败：{e}")


@app.get("/api/ledger")
def ledger():
    cfg = get_cfg()
    data = zjb.load_ledger(cfg)
    return {"ledger_file": str(zjb.ledger_file(cfg)), "weeks": data.get("weeks", {})}


@app.post("/api/sync-assistants")
def sync_assistants(body: SyncAssistIn):
    """联动：把 05 工作安排台某周的助理安排，写回 04 台账对应周次活动的「检查助理」。"""
    cfg = get_cfg()
    return do_sync_assistants(cfg, body.source_week, body.target_week, body.task, body.plan_url)


@app.get("/api/plan-meta")
def plan_meta(plan_url: str = "http://127.0.0.1:8766/api/plan/weeks"):
    """联动面板用：问 05 工作安排台有哪些周的计划、各任务叫什么，供下拉选择（服务端代理，免跨域）。"""
    try:
        with urllib.request.urlopen(plan_url, timeout=6) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        return {"ok": True, "weeks": data.get("weeks", []), "tasks": data.get("tasks", {})}
    except Exception as e:
        return {"ok": False, "error": str(e), "weeks": [], "tasks": {}}


@app.delete("/api/ledger")
def ledger_reset(week: Optional[int] = None, campus: Optional[str] = None):
    cfg = get_cfg()
    ledger = zjb.load_ledger(cfg)
    weeks = ledger.get("weeks", {})
    if week and campus:
        weeks.pop(f"W{week}-{campus}", None)
    elif week:
        for k in [k for k in weeks if k.startswith(f"W{week}-")]:
            weeks.pop(k, None)
    else:
        weeks.clear()
    zjb.save_ledger(cfg, ledger)
    return {"ok": True, "remaining": len(weeks)}


@app.get("/api/stats")
def stats():
    return build_stats(get_cfg())


@app.get("/api/history")
def history():
    """已上传/生成的历史文件清单，按周分组，供前端「历史归档」浏览与下载。

    数据来自 registry.json（每次上传、拆分、安排、助理都会登记），
    含 path/week/campus/kind/created_at；这里补全 download_url 与文件是否存在。"""
    reg = load_registry()
    groups: Dict[int, List[dict]] = {}
    for fid, item in reg.items():
        w = item.get("week")
        if w is None:
            continue
        groups.setdefault(int(w), []).append({
            "id": fid,
            "name": item.get("name"),
            "kind": item.get("kind"),
            "campus": item.get("campus"),
            "week": w,
            "created_at": item.get("created_at"),
            "download_url": f"/api/files/{fid}/download",
            "exists": Path(item.get("path", "")).exists(),
        })
    out = []
    for w in sorted(groups.keys()):
        files = sorted(groups[w], key=lambda x: x.get("created_at") or "")
        out.append({"week": w, "files": files})
    return {"weeks": out, "total": len(reg)}


@app.get("/api/slots/template")
def slots_template():
    return {"周六上午": [], "周六下午": [], "周日上午": [], "周日下午": []}


# ---------------------------------------------------------------- 预览与体检

def _rgb(v, default="9aa0a6"):
    try:
        s = getattr(v, "rgb", None)
        if isinstance(s, str) and len(s) >= 6:
            return s[-6:]
    except Exception:
        pass
    return default


def _last_used_col(ws, max_scan=None) -> int:
    """最后一个真正有内容或边框的列（避开 openpyxl 的幽灵列）"""
    limit = max_scan or min(ws.max_column or 1, 60)
    last = 1
    for c in range(limit, 0, -1):
        for r in range(1, min(ws.max_row or 1, 200) + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is not None and str(cell.value).strip():
                return c
            b = cell.border
            if b and any(getattr(b, s) and getattr(b, s).style for s in ("left", "right", "top", "bottom")):
                last = max(last, c)
                break
    return last


def render_builtin_html(src: Path) -> str:
    """内置渲染回退：用 openpyxl 读样式，还原字体、字号、加粗、对齐、边框、合并单元格"""
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    ncol = _last_used_col(ws)
    nrow = ws.max_row or 1

    merged = {}
    covered = set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1,
                                              rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))

    def cell_style(cell):
        css = []
        f = cell.font
        if f:
            if f.name:
                css.append(f"font-family:'{f.name}',SimSun,serif")
            if f.size:
                css.append(f"font-size:{f.size}pt")
            if f.bold:
                css.append("font-weight:700")
        a = cell.alignment
        if a:
            if a.horizontal:
                css.append(f"text-align:{a.horizontal}")
            if a.vertical:
                css.append(f"vertical-align:{'middle' if a.vertical == 'center' else a.vertical}")
            if a.wrap_text:
                css.append("white-space:pre-wrap")
        b = cell.border
        if b:
            for side in ("left", "right", "top", "bottom"):
                s = getattr(b, side)
                if s and s.style:
                    css.append(f"border-{side}:1px solid #{_rgb(getattr(s, 'color', None))}")
        fill = cell.fill
        if fill and getattr(fill, "patternType", None):
            css.append(f"background:#{_rgb(getattr(fill, 'fgColor', None))}")
        return ";".join(css)

    rows_html = []
    for r in range(1, nrow + 1):
        tds = []
        for c in range(1, ncol + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            rowspan, colspan = merged.get((r, c), (1, 1))
            val = "" if cell.value is None else str(cell.value)
            val = (val.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
            span = ""
            if rowspan > 1:
                span += f' rowspan="{rowspan}"'
            if colspan > 1:
                span += f' colspan="{colspan}"'
            tds.append(f'<td{span} style="{cell_style(cell)}">{val}</td>')
        if tds:
            rows_html.append("<tr>" + "".join(tds) + "</tr>")

    title = src.name.replace("&", "&amp;").replace("<", "&lt;")
    return (
        "<!DOCTYPE html><html><head><meta charset='UTF-8'>"
        f"<title>{title}</title><style>"
        "body{margin:0;padding:14px;background:#fff;font-family:'SimSun',serif;}"
        "table{border-collapse:collapse;table-layout:fixed;}"
        "td{border:1px solid #d7dbe0;padding:3px 6px;font-size:9pt;line-height:1.4;"
        "word-break:break-all;}"
        ".hint{font-size:12px;color:#6b7480;margin-bottom:8px;}"
        "</style></head><body>"
        "<div class='hint'>内置渲染（OfficeCLI 不可用时的回退视图）· " + title + "</div>"
        "<table>" + "".join(rows_html) + "</table></body></html>"
    )


def render_preview(src: Path) -> dict:
    """生成预览 HTML：优先 OfficeCLI（高保真），不可用时回退内置渲染"""
    key = hashlib.md5(str(src).encode("utf-8")).hexdigest()[:16]
    out_html = PREVIEW_DIR / f"{key}.html"
    stamp = src.stat().st_mtime if src.exists() else 0
    meta_file = PREVIEW_DIR / f"{key}.json"
    fresh = False
    if meta_file.exists() and out_html.exists():
        try:
            meta = json.loads(meta_file.read_text(encoding="utf-8"))
            fresh = meta.get("mtime") == stamp and meta.get("engine") is not None
        except Exception:
            fresh = False

    engine = "builtin"
    if not fresh:
        if OFFICECLI.exists():
            try:
                proc = subprocess.run(
                    [str(OFFICECLI), "view", str(src), "html", "-o", str(out_html)],
                    capture_output=True, timeout=120)
                if proc.returncode == 0 and out_html.exists():
                    engine = "officecli"
            except Exception:
                engine = "builtin"
        if engine == "builtin":
            out_html.write_text(render_builtin_html(src), encoding="utf-8")
        meta_file.write_text(json.dumps({"engine": engine, "mtime": stamp}, ensure_ascii=False),
                             encoding="utf-8")
    else:
        try:
            engine = json.loads(meta_file.read_text(encoding="utf-8")).get("engine", "builtin")
        except Exception:
            engine = "builtin"
    return {"url": f"/preview/{out_html.name}", "engine": engine, "name": src.name}


@app.post("/api/preview")
def preview(body: SimpleIn):
    src = resolve(body.file_id)
    return render_preview(src)


@app.get("/api/preview/file/{fid}")
def preview_by_id(fid: str):
    return render_preview(resolve(fid))


@app.post("/api/diagnose")
def diagnose(body: SimpleIn):
    """下载前体检：可打开性 + 结构完整性 + 业务列完整性 + OpenXML 校验"""
    cfg = get_cfg()
    src = resolve(body.file_id)
    checks: List[Dict[str, Any]] = []

    try:
        wb = openpyxl.load_workbook(src)
        ws = wb.active
        checks.append({"name": "文件可正常打开", "ok": True, "detail": f"{ws.max_row} 行 × {ws.max_column} 列"})
    except Exception as e:
        return {"ok": False, "file": src.name,
                "checks": [{"name": "文件可正常打开", "ok": False, "detail": str(e)}]}

    header_row, header_map = zjb.find_header_row(ws)
    checks.append({"name": "表头识别", "ok": bool(header_map),
                   "detail": f"第 {header_row} 行，识别到 {len(header_map)} 个标准列"})

    need = {"name": "活动名称", "host": "主办单位", "campus": "校区", "time": "开展时间"}
    missing = [label for k, label in need.items() if not header_map.get(k)]
    checks.append({"name": "关键列齐全", "ok": not missing,
                   "detail": "全部具备" if not missing else "缺少：" + "、".join(missing)})

    merged = [str(r) for r in ws.merged_cells.ranges]
    title_merge = next((m for m in merged if m.startswith("A1")), None)
    checks.append({"name": "标题合并区", "ok": bool(title_merge),
                   "detail": title_merge or "未发现 A1 起始的合并区（标题可能丢失）"})

    acts = zjb.read_activities(ws, header_row, header_map, cfg)
    checks.append({"name": "活动数据", "ok": len(acts) > 0,
                   "detail": f"共 {len(acts)} 个活动" if acts else "未读到任何活动行"})

    _, checker_map = zjb.read_checkers(ws, header_row, cfg.get("checker_column", []), cfg)
    per = cfg.get("per_activity_colleges", 2)
    if acts:
        filled = sum(1 for a in acts if checker_map.get(a["row"]))
        bad = [a["name"] for a in acts if len(checker_map.get(a["row"], [])) != per]
        ok_cnt = not bad
        checks.append({"name": "检查学院安排", "ok": ok_cnt and filled == len(acts),
                       "detail": f"{filled}/{len(acts)} 个活动已安排" +
                                 ("" if ok_cnt else "，数量不是 %d 个：" % per + "、".join(bad[:3]))})

    if OFFICECLI.exists():
        try:
            proc = subprocess.run([str(OFFICECLI), "validate", str(src)],
                                  capture_output=True, timeout=120)
            out = (proc.stdout or b"").decode("utf-8", "replace").strip()
            ok = proc.returncode == 0
            checks.append({"name": "OpenXML 结构校验（OfficeCLI）", "ok": ok,
                           "detail": (out.splitlines() or [""])[0][:200] or ("通过" if ok else "未通过")})
        except Exception as e:
            checks.append({"name": "OpenXML 结构校验（OfficeCLI）", "ok": True,
                           "detail": f"跳过（{type(e).__name__}）"})
    else:
        checks.append({"name": "OpenXML 结构校验", "ok": True, "detail": "跳过（未安装 OfficeCLI）"})

    return {"ok": all(c["ok"] for c in checks), "file": src.name, "checks": checks}


@app.get("/api/files/{fid}/download")
def download(fid: str):
    src = resolve(fid)
    return FileResponse(str(src), filename=src.name,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class NoCacheStatic(StaticFiles):
    """静态文件禁用浏览器缓存：重新构建前端后刷新页面即可生效。

    否则浏览器会按 etag/Last-Modified 复用旧 JS，用户会以为「前端没变」。
    """

    def file_response(self, *args, **kwargs):
        resp = super().file_response(*args, **kwargs)
        resp.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        resp.headers["Pragma"] = "no-cache"
        resp.headers["Expires"] = "0"
        return resp


# 预览产物（OfficeCLI / 内置渲染生成的独立 HTML）
app.mount("/preview", NoCacheStatic(directory=str(PREVIEW_DIR), html=True), name="preview")

# 落地页已并入工作台首页（React 的 / 首屏就是「档案台」hero）——
# 这里做兼容跳转：老链接/书签里的 /landing 与 /landing/ 都 302 到 /。
# 早期版本用 StaticFiles 挂载，踩过坑：/landing（无尾斜杠）会 404，
# 而启动脚本恰恰打开的就是这个地址 → 用户看到空白页。改跳转后彻底消除。
@app.get("/landing")
@app.get("/landing/{rest:path}")
def landing_redirect():
    return RedirectResponse("/", status_code=302)

# 静态前端（生产构建产物）
if WEB_DIST.exists():
    app.mount("/", NoCacheStatic(directory=str(WEB_DIST), html=True), name="web")
else:
    @app.get("/")
    def dev_hint():
        return HTMLResponse(
            "<h3>前端尚未构建</h3><p>请先执行 <code>cd web && npm run build</code>，"
            "或用开发模式 <code>npm run dev</code> 访问 5173 端口。</p>")


if __name__ == "__main__":
    import uvicorn
    # 发布沙箱注入 PORT 并只暴露 0.0.0.0；本地仍可用 ZJB_PORT 覆盖
    uvicorn.run(app, host="0.0.0.0",
                port=int(os.environ.get("PORT", os.environ.get("ZJB_PORT", "8765"))),
                log_level="warning")
